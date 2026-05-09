# -*- coding: utf-8 -*-
"""
职业转换卡系统策划案生成器
"""

import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
sys.stdout.reconfigure(encoding='utf-8')

def create_doc(output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = '职业转换卡系统策划案'

    # 样式定义
    title_font   = Font(name='微软雅黑', size=14, bold=True, color='FFFFFF')
    h1_font      = Font(name='微软雅黑', size=12, bold=True)
    h2_font      = Font(name='微软雅黑', size=11, bold=True)
    h3_font      = Font(name='微软雅黑', size=10, bold=True)
    body_font    = Font(name='微软雅黑', size=10)

    title_fill   = PatternFill(start_color='2E4057', end_color='2E4057', fill_type='solid')
    h1_fill      = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    h2_fill      = PatternFill(start_color='B4C7E7', end_color='B4C7E7', fill_type='solid')
    h3_fill      = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    warn_fill    = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

    # 列宽
    ws.column_dimensions['A'].width = 4
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['C'].width = 70
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 18

    r = 1  # 当前行

    def write(row, col, val, font=None, fill=None, bold=False, indent=0, merge_to=None, wrap=True, height=22):
        prefix = '　' * indent
        cell = ws.cell(row=row, column=col, value=prefix + str(val) if val else val)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(vertical='center', wrap_text=wrap, indent=0)
        ws.row_dimensions[row].height = height
        if merge_to:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=merge_to)
        return row + 1

    def blank(row, n=1):
        for _ in range(n):
            ws.row_dimensions[row].height = 8
            row += 1
        return row

    def title_row(row, text):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = title_font
        cell.fill = title_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[row].height = 30
        return row + 1

    def h1(row, text):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
        cell.fill = h1_fill
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[row].height = 26
        return row + 1

    def h2(row, text):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = h2_font
        cell.fill = h2_fill
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[row].height = 22
        return row + 1

    def h3(row, text):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=2, value=text)
        cell.font = h3_font
        cell.fill = h3_fill
        cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[row].height = 20
        return row + 1

    def row2(row, label, content, label_fill=None, content_fill=None, height=20):
        cell_l = ws.cell(row=row, column=2, value=label)
        cell_l.font = h3_font
        cell_l.alignment = Alignment(vertical='center', wrap_text=True)
        if label_fill:
            cell_l.fill = label_fill
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
        cell_r = ws.cell(row=row, column=3, value=content)
        cell_r.font = body_font
        cell_r.alignment = Alignment(vertical='center', wrap_text=True)
        if content_fill:
            cell_r.fill = content_fill
        ws.row_dimensions[row].height = height
        return row + 1

    def bullet(row, text, level=1, fill=None, height=20):
        prefix = {1: '·', 2: '  - ', 3: '    · '}
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=2, value=prefix.get(level, '') + text)
        cell.font = h3_font if level == 1 else body_font
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if fill:
            cell.fill = fill
        ws.row_dimensions[row].height = height
        return row + 1

    # ===================== 封面 =====================
    r = title_row(r, '职业转换卡系统 · 系统策划案')
    r = blank(r)
    r = row2(r, '功能名称', '职业转换卡系统')
    r = row2(r, '文档日期', '')
    r = row2(r, '操作者', '')
    r = row2(r, '版本', 'V1.0')
    r = row2(r, '备注', '')
    r = blank(r, 2)

    # ===================== 一、系统介绍 =====================
    r = h1(r, '一、系统介绍')
    r = h2(r, '1.1 功能说明')
    r = bullet(r, '职业转换卡是一种消耗道具，使用后可将角色转换为指定的非本职业的职业，转换为无损转换。')
    r = bullet(r, '转换卡分为两类：', 1)
    r = bullet(r, '指定职业无损转换卡：只能转换为卡片指定的职业', 2)
    r = bullet(r, '任选职业无损转换卡：可自由选择任意非本职业进行转换', 2)
    r = blank(r)

    r = h2(r, '1.2 功能入口')
    r = bullet(r, '入口位置：角色界面 → 新增页签"职业转换"（放在最后面）')
    r = bullet(r, '背包使用：在背包中点击转换卡，跳转至职业转换界面')
    r = blank(r)

    r = h2(r, '1.3 功能开关')
    r = bullet(r, '功能开关：运营后台配置，关闭时入口隐藏，转换卡不可使用')
    r = blank(r, 2)

    # ===================== 二、使用前置条件 =====================
    r = h1(r, '二、使用前置条件')
    r = bullet(r, '以下所有条件须同时满足，任一不满足则飘字提示对应原因：', 1)
    conditions = [
        ('可转换的职业', '转换卡配置了可转换的目标职业，且目标职业非本职业'),
        ('可使用等级', '角色等级达到配置的最低使用等级'),
        ('开服天数限制', '当前开服天数满足配置的最低天数要求'),
        ('功能开关', '功能开关处于开启状态'),
        ('非转职冷却时间', '距上次转职已超过冷却时间'),
        ('非战斗状态', '角色当前不在战斗中'),
        ('非本职业', '目标职业与当前职业不同'),
        ('在主城', '角色当前位于主城场景'),
        ('消耗道具足够', '背包中有足够的转换所需消耗道具'),
    ]
    for cond, desc in conditions:
        r = row2(r, cond, desc, height=22)
    r = blank(r, 2)

    # ===================== 三、职业转换界面 =====================
    r = h1(r, '三、职业转换界面')
    r = h2(r, '3.1 展示信息')
    display_items = [
        '所有职业列表（本职业置底，灰化不可选）',
        '各职业所有主动技能列表',
        '技能简易描述',
        '各职业模型展示（用于展示技能效果）',
        '转职规则说明',
        '转职消耗道具及数量',
        '转职按钮',
    ]
    for item in display_items:
        r = bullet(r, item, 2)
    r = blank(r)

    r = h2(r, '3.2 转职操作流程')
    r = bullet(r, '步骤1：玩家在职业转换界面选择目标职业')
    r = bullet(r, '步骤2：点击"转职"按钮，系统判断所有前置条件')
    r = bullet(r, '步骤3a：条件满足 → 消耗道具 → 弹出二次确认弹窗')
    r = bullet(r, '步骤3b：条件不满足 → 对应飘字提示，不消耗道具')
    r = bullet(r, '步骤4：玩家确认后，进入转职进度界面（期间不可进行任何操作）')
    r = bullet(r, '步骤5：转职成功 → 弹窗提示"转职成功，需要重启游戏"')
    r = bullet(r, '步骤6：玩家关闭弹窗 → 游戏重启')
    r = blank(r)

    r = h2(r, '3.3 转职中断处理')
    r = bullet(r, '转职期间关闭游戏：')
    r = bullet(r, '重新登录转职服务器时，若转职未完成 → 继续弹出转职进行中界面', 2)
    r = bullet(r, '重新登录转职服务器时，若转职已完成 → 按正常流程进入游戏', 2)
    r = blank(r, 2)

    # ===================== 四、转换细则 =====================
    r = h1(r, '四、转换细则（无损继承规则）')

    # 角色基础
    r = h2(r, '4.1 角色基础属性')
    items_41 = [
        ('等级、转职、转生', '无损继承'),
        ('性别', '无损继承'),
        ('VIP', '无损继承'),
        ('所有货币', '无损继承'),
        ('经验进度', '保持不变'),
        ('结婚伴侣', '保持不变'),
        ('公会', '保持不变'),
    ]
    for k, v in items_41:
        r = row2(r, k, v)
    r = blank(r)

    # 技能
    r = h2(r, '4.2 技能')
    items_42 = [
        ('技能', '转换为新职业对应技能'),
        ('技能升级/突破/觉醒等级', '继承（等级数值不变，映射到新职业对应技能）'),
        ('技能配置', '保存并自动适配新职业对应技能'),
        ('技能书/突破道具/觉醒道具', '转换为新职业对应道具'),
        ('被动技能', '保持不变'),
    ]
    for k, v in items_42:
        r = row2(r, k, v)
    r = blank(r)

    # 时装
    r = h2(r, '4.3 时装')
    items_43 = [
        ('时装激活', '自动激活新职业对应时装，继承星级'),
        ('时装道具', '本职业时装道具转换为新职业对应时装道具'),
        ('时装染色', '继承'),
        ('化妆相关', '继承'),
    ]
    for k, v in items_43:
        r = row2(r, k, v)
    r = blank(r)

    # 装备
    r = h2(r, '4.4 装备（穿戴中/背包/仓库/邮件/装备收藏）')
    items_44 = [
        ('装备本体', '本职业装备转换为新职业对应装备（绑定状态/品质/神装/祝福不变）'),
        ('魔戒', '不区分职业，不特殊处理'),
        ('强化', '继承'),
        ('镶嵌', '继承'),
        ('追加', '继承'),
        ('卡片', '继承'),
        ('祝福', '继承'),
        ('套装激活', '继承；套装石转换为新职业对应套装石；套装回退时返还新职业对应套装石'),
        ('符文', '继承；本职业符文转换为新职业对应符文'),
        ('神饰', '继承'),
        ('装备收藏', '收藏的装备转换为对应职业装备，再生属性无损继承'),
    ]
    for k, v in items_44:
        r = row2(r, k, v, height=24)
    r = blank(r)

    # 背包/仓库/邮件/交易行道具
    r = h2(r, '4.5 背包/仓库/邮件/交易行道具')
    items_45 = [
        ('装备类', '本职业装备转换为新职业对应装备，绑定状态不变'),
        ('职业相关道具', '技能书、技能突破/觉醒、套装石碎片/宝箱/套装石、神兵/神兵碎片、装备宝箱等均转换为新职业对应道具'),
        ('非职业相关道具', '无损继承，包括数量和绑定状态'),
        ('活动临时仓库', '按背包/邮件道具规则直接转化处理'),
        ('交易行上架商品', '保持原上架状态和时间；本职业装备/道具转换为新职业对应；已关注道具自动刷新为新职业相关，无需玩家重新关注'),
        ('商店/购物单', '与本职业相关的道具转换为新职业对应道具'),
    ]
    for k, v in items_45:
        r = row2(r, k, v, height=30)
    r = blank(r)

    # 神兵
    r = h2(r, '4.6 神兵')
    items_46 = [
        ('已激活神兵', '转职后对应新职业神兵自动激活，继承升星星级（需配置一对一对应关系）'),
        ('穿戴/背包/仓库/邮件/临时仓库神兵', '转换为新职业对应神兵'),
        ('神兵碎片', '转换为新职业对应碎片'),
        ('器灵', '无损继承'),
        ('器灵触发技能', '根据神兵替换为对应技能'),
    ]
    for k, v in items_46:
        r = row2(r, k, v, height=24)
    r = blank(r)

    # 其他系统
    r = h2(r, '4.7 其他系统')
    items_47 = [
        ('称号', '无损继承'),
        ('装扮', '星级无损继承'),
        ('头衔', '等级无损继承'),
        ('转生之证', '转生点、技能等级无损继承'),
        ('神馈', '保持不变'),
        ('吞噬', '等级、特权无损继承'),
        ('英雄', '阵容/圣坛/英雄等级/洗髓/星级/英雄装备/英雄戒指/英雄魂石/圣树/UR/幻化/图鉴/众神谱无损继承'),
        ('伙伴（坐骑/翅膀/守护）', '无损继承'),
        ('助战（神像/契约/泰坦/奖杯/圣物）', '无损继承'),
        ('神佑（英雄圣装/神盾/神甲/神纹）', '无损继承'),
        ('伴侣', '无损继承'),
        ('合成', '显示新职业合成配方'),
    ]
    for k, v in items_47:
        r = row2(r, k, v, height=22)
    r = blank(r)

    # 外观与奖励
    r = h2(r, '4.8 外观与奖励匹配')
    items_48 = [
        ('新职业外观', '主城雕像/排行榜/天神骑士团/位面霸主/竞技场结算/王者1v1/结缘等显示匹配新职业'),
        ('奖励匹配', '福利/直购/运营活动/BP/寻宝/首充等与职业相关的奖励匹配新职业'),
        ('转职称号/成就', '对应转职的职业称号、成就匹配并继承对应转职阶段'),
        ('进行中的转职', '刷新为对应新职业的转职进度'),
        ('穿戴外显', '保留玩家原本穿戴模型（装备/神兵/时装等外显）'),
    ]
    for k, v in items_48:
        r = row2(r, k, v, height=28)
    r = blank(r, 2)

    # ===================== 五、异常处理 =====================
    r = h1(r, '五、异常处理')
    r = h2(r, '5.1 转职流程异常')
    r = bullet(r, '转职中途断线：重连后继续弹出转职进行中界面，直至转职完成')
    r = bullet(r, '转职完成后断线：下次登录按正常流程进入游戏，无需额外处理')
    r = bullet(r, '道具消耗成功但转职失败：服务端回滚道具，客户端提示"转职失败，道具已返还"')
    r = bullet(r, '转职期间被踢下线：重新登录后检查转职状态，未完成则继续，已完成则正常进入')
    r = blank(r)

    r = h2(r, '5.2 数据转换异常')
    r = bullet(r, '职业对应关系未配置：转换时跳过该类道具，记录日志，GM后台可手动补偿')
    r = bullet(r, '目标职业道具不存在：保留原道具，记录异常日志，告警通知策划/程序')
    r = bullet(r, '数量溢出：超出上限部分转入邮件，附带说明')
    r = bullet(r, '背包空间不足：优先放入背包，不足部分转入邮件')
    r = blank(r)

    r = h2(r, '5.3 操作异常')
    r = bullet(r, '重复点击转职按钮：按钮点击后进入loading状态，转职完成前不可再次点击')
    r = bullet(r, '转职期间切换账号：强制中断转职，道具回滚，提示"转职已取消"')
    r = bullet(r, '转职期间游戏崩溃：重启后检查转职状态，按5.1规则处理')
    r = blank(r, 2)

    # ===================== 六、边界条件 =====================
    r = h1(r, '六、边界条件')
    r = bullet(r, '同一账号同时在多台设备操作转职：后操作设备提示"当前账号正在转职中，请稍后"')
    r = bullet(r, '转职冷却时间内再次尝试：飘字提示剩余冷却时间')
    r = bullet(r, '转职目标职业与当前职业相同：按钮置灰，不可点击')
    r = bullet(r, '背包/仓库/邮件中道具数量为0：跳过转换，不报错')
    r = bullet(r, '转换卡配置的目标职业不存在：转换卡不可使用，提示"该职业暂未开放"')
    r = bullet(r, '开服天数不足：飘字提示"开服第X天后可使用"')
    r = blank(r, 2)

    # ===================== 七、测试用例 =====================
    r = h1(r, '七、测试用例')
    r = h2(r, '7.1 功能测试')
    test_cases = [
        ('TC-001', '指定职业转换卡正常转职', '满足所有前置条件，使用指定职业转换卡', '转职成功，数据按规则转换，弹窗提示重启'),
        ('TC-002', '任选职业转换卡正常转职', '满足所有前置条件，使用任选职业转换卡选择目标职业', '转职成功，数据按规则转换'),
        ('TC-003', '前置条件不满足-战斗中', '角色在战斗中点击转职', '飘字提示"请在非战斗状态下使用"'),
        ('TC-004', '前置条件不满足-非主城', '角色在副本中点击转职', '飘字提示"请在主城使用"'),
        ('TC-005', '前置条件不满足-冷却中', '转职冷却时间内再次转职', '飘字提示剩余冷却时间'),
        ('TC-006', '转职中断-关闭游戏', '转职进行中关闭游戏，重新登录', '继续弹出转职进行中界面'),
        ('TC-007', '转职完成后重启', '转职成功后重启游戏', '正常进入游戏，职业已变更'),
        ('TC-008', '装备转换验证', '转职后检查穿戴装备、背包装备', '所有装备转换为新职业对应装备，属性/绑定状态不变'),
        ('TC-009', '技能继承验证', '转职后检查技能等级/突破/觉醒', '等级数值继承，技能映射到新职业'),
        ('TC-010', '神兵转换验证', '转职后检查神兵激活状态和碎片', '对应神兵激活，碎片转换正确'),
    ]
    # 写表头
    headers = ['用例编号', '用例名称', '操作步骤', '预期结果']
    for col_idx, h in enumerate(headers, start=2):
        cell = ws.cell(row=r, column=col_idx, value=h)
        cell.font = h3_font
        cell.fill = h3_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[r].height = 22
    r += 1
    for tc in test_cases:
        for col_idx, val in enumerate(tc, start=2):
            cell = ws.cell(row=r, column=col_idx, value=val)
            cell.font = body_font
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 24
        r += 1
    r = blank(r)

    r = h2(r, '7.2 暴力测试')
    r = bullet(r, '快速连续点击转职按钮100次：不产生重复转职，按钮正确进入loading状态')
    r = bullet(r, '转职进行中快速切换界面：转职流程不中断，界面切回后状态正确')
    r = bullet(r, '背包满载时转职：道具转换后溢出部分正确进入邮件')
    r = bullet(r, '同时持有大量职业相关道具（1000+）时转职：全部正确转换，无遗漏')
    r = bullet(r, '网络极差（丢包率50%）时转职：转职流程最终完成或正确回滚')
    r = blank(r)

    r = h2(r, '7.3 验收标准')
    r = bullet(r, '所有TC用例通过率100%')
    r = bullet(r, '无P0/P1级别Bug')
    r = bullet(r, '转职数据转换准确率100%（抽查100个账号）')
    r = bullet(r, '转职流程平均耗时小于等于30秒')
    r = blank(r, 2)

    # ===================== 八、埋点需求 =====================
    r = h1(r, '八、埋点需求')
    r = h2(r, '8.1 曝光埋点')
    r = bullet(r, '职业转换界面曝光：次数、人数、来源入口（角色页签/背包）')
    r = bullet(r, '各职业展示曝光：每个职业被查看的次数')
    r = blank(r)

    r = h2(r, '8.2 点击埋点')
    r = bullet(r, '转职按钮点击：次数、人数、目标职业')
    r = bullet(r, '二次确认弹窗-确认：次数、人数')
    r = bullet(r, '二次确认弹窗-取消：次数、人数')
    r = blank(r)

    r = h2(r, '8.3 结果埋点')
    r = bullet(r, '转职成功：次数、人数、来源职业、目标职业、使用的转换卡类型')
    r = bullet(r, '转职失败：次数、失败原因分布（前置条件不满足类型）')
    r = bullet(r, '转换卡消耗：消耗数量、人数、转换卡类型')
    r = blank(r)

    r = h2(r, '8.4 性能埋点')
    r = bullet(r, '转职流程耗时：从确认到完成的时长分布')
    r = bullet(r, '转职界面加载时长')
    r = blank(r, 2)

    # ===================== 九、功能开关与数值配置 =====================
    r = h1(r, '九、功能开关与数值配置')
    config_items = [
        ('功能总开关', '布尔值，控制整个职业转换功能的开启/关闭'),
        ('最低使用等级', '整数，角色需达到的最低等级'),
        ('开服天数限制', '整数，开服后多少天才可使用'),
        ('转职冷却时间', '整数（秒），两次转职之间的最短间隔'),
        ('指定职业转换卡配置', '道具ID → 可转换职业ID列表'),
        ('任选职业转换卡配置', '道具ID → 可转换职业ID列表（全职业）'),
        ('职业对应关系表', '原职业ID → 新职业ID，用于装备/道具/神兵等转换映射'),
    ]
    for k, v in config_items:
        r = row2(r, k, v, height=24)
    r = blank(r, 2)

    # 保存
    wb.save(output_path)
    print(f'策划案已生成: {output_path}')

if __name__ == '__main__':
    output = r'E:\桌面\其它\职业转换卡系统策划案.xlsx'
    create_doc(output)
