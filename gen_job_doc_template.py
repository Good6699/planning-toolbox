# -*- coding: utf-8 -*-
"""
职业转换卡系统策划案 - 使用原模版格式
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import sys
sys.stdout.reconfigure(encoding='utf-8')

def create_doc():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    
    # 样式定义
    normal_font = Font(name='微软雅黑', size=10)
    bold_font = Font(name='微软雅黑', size=10, bold=True)
    
    # 设置列宽（匹配原模版）
    ws.column_dimensions['A'].width = 4.25
    ws.column_dimensions['B'].width = 4.5
    ws.column_dimensions['C'].width = 4.5
    ws.column_dimensions['D'].width = 4.25
    ws.column_dimensions['E'].width = 13.0
    ws.column_dimensions['F'].width = 13.0
    ws.column_dimensions['G'].width = 13.0
    ws.column_dimensions['H'].width = 13.0
    ws.column_dimensions['I'].width = 13.0
    ws.column_dimensions['J'].width = 13.0
    ws.column_dimensions['K'].width = 13.0
    ws.column_dimensions['L'].width = 13.0
    ws.column_dimensions['M'].width = 2.875
    ws.column_dimensions['N'].width = 4.25
    ws.column_dimensions['O'].width = 13.0
    ws.column_dimensions['P'].width = 13.0
    ws.column_dimensions['Q'].width = 13.0
    ws.column_dimensions['R'].width = 13.0
    ws.column_dimensions['S'].width = 13.0
    ws.column_dimensions['T'].width = 13.0
    ws.column_dimensions['U'].width = 13.0
    ws.column_dimensions['V'].width = 13.0
    ws.column_dimensions['W'].width = 13.0
    ws.column_dimensions['X'].width = 13.0
    ws.column_dimensions['Y'].width = 13.0
    ws.column_dimensions['Z'].width = 13.0
    ws.column_dimensions['AA'].width = 13.0
    ws.column_dimensions['AB'].width = 13.0
    ws.column_dimensions['AC'].width = 13.0
    ws.column_dimensions['AD'].width = 13.0
    ws.column_dimensions['AE'].width = 13.0
    
    # 合并单元格（日期/操作者区域）
    ws.merge_cells('B3:E3')
    ws.merge_cells('B4:E4')
    ws.merge_cells('B5:E5')
    ws.merge_cells('B6:E6')
    ws.merge_cells('F3:H3')
    ws.merge_cells('F4:H4')
    ws.merge_cells('F5:H5')
    ws.merge_cells('F6:H6')
    ws.merge_cells('I3:N3')
    ws.merge_cells('I4:N4')
    ws.merge_cells('I5:N5')
    ws.merge_cells('I6:N6')
    
    # 基础信息区域
    ws['B1'] = '功能名称'
    ws['B3'] = '日期'
    ws['F3'] = '操作者'
    ws['I3'] = '备注'
    ws['F4'] = '赵伟'
    ws['I4'] = '建立'
    
    r = 10  # 系统介绍从第10行开始（原模版结构）
    
    # 系统介绍
    ws.cell(row=r, column=2, value='系统介绍')
    r += 2
    
    ws.cell(row=r, column=2, value='·功能说明')
    r += 1
    ws.cell(row=r, column=3, value='·职业转换卡是一种消耗道具，使用后可将角色转换为指定的非本职业的职业，转换为无损转换')
    r += 1
    ws.cell(row=r, column=3, value='·转换卡分为两类：指定职业无损转换卡、任选职业无损转换卡')
    r += 2
    
    ws.cell(row=r, column=2, value='·功能要求')
    r += 1
    ws.cell(row=r, column=3, value='·功能开启条件：可转换职业、可使用等级、开服天数限制、功能开关开启、非转职冷却时间、非战斗状态、非本职业、在主城、消耗道具足够')
    r += 2
    
    ws.cell(row=r, column=2, value='·活动流程')
    r += 1
    ws.cell(row=r, column=3, value='·背包里使用：点击跳转至职业转换界面')
    r += 1
    ws.cell(row=r, column=3, value='·职业转换界面入口：放在角色里面，新加页签"职业转换"，放在最后面')
    r += 1
    ws.cell(row=r, column=3, value='·展示信息：所有职业（本职业置底，灰化）、各职业所有技能（只展示主动技能）、技能简易描述、各职业模型、转职规则、转职消耗、转职按钮')
    r += 1
    ws.cell(row=r, column=3, value='·点击转职按钮后判断前置条件是否满足，都满足后消耗道具，弹窗2次确认')
    r += 1
    ws.cell(row=r, column=3, value='·确认后开始展示转职进度界面，期间玩家不能进行任何操作')
    r += 1
    ws.cell(row=r, column=3, value='·转职成功后弹窗提醒"转职成功，需要重启游戏"，关闭弹窗重启游戏')
    r += 1
    ws.cell(row=r, column=3, value='·转职期间关掉游戏后，登录转职的服务器时如果转职未完成，则继续弹转职进行中界面')
    r += 1
    ws.cell(row=r, column=3, value='·如果转职已完成，则按照正常流程进入游戏即可')
    r += 2
    
    ws.cell(row=r, column=2, value='·功能点名称')
    r += 1
    ws.cell(row=r, column=3, value='·角色：等级、转职、转生、性别、VIP、所有货币无损继承；结婚伴侣、公会保持不变；经验进度保持不变')
    r += 1
    ws.cell(row=r, column=3, value='·神馈：保持不变')
    r += 1
    ws.cell(row=r, column=3, value='·技能：转换为新职业对应技能，继承技能升级、突破、觉醒等级，技能配置保存自动适配新职业，本职业技能书/突破道具/觉醒道具转换为新职业的，被动保持不变')
    r += 1
    ws.cell(row=r, column=3, value='·时装：自动激活新职业对应时装且继承星级，本职业时装道具转换为新职业对应时装道具，时装染色继承，化妆相关修改继承')
    r += 1
    ws.cell(row=r, column=3, value='·称号：无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·装扮：星级无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·头衔：等级无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·转生之证：转生点、技能等级无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·背包/仓库/邮件道具/交易行：装备类本职业装备转换为新职业对应装备，绑定状态不变；道具类本职业道具转换为新职业对应道具；其它与本职业非关的道具无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·活动临时仓库：将临时仓库中的道具按照背包、邮件中道具规则直接转化处理')
    r += 1
    ws.cell(row=r, column=3, value='·合成：显示新职业合成配方')
    r += 1
    ws.cell(row=r, column=3, value='·吞噬：等级、特权无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·英雄：阵容、圣坛、英雄等级、洗髓、星级、英雄装备、英雄戒指、英雄魂石、圣树、UR、幻化、图鉴、众神谱无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·装备转换（穿戴中/背包/仓库/邮件/装备收藏）：原职业穿戴的装备转化为新职业对应装备，绑定状态不变、品质不变、神装不变、祝福不变')
    r += 1
    ws.cell(row=r, column=3, value='·魔戒：魔戒不区分职业，不特殊处理')
    r += 1
    ws.cell(row=r, column=3, value='·装备锻造：强化、镶嵌、追加、卡片、祝福、套装激活、符文、神饰需继承，本职业套装石/符文转换为新职业对应的')
    r += 1
    ws.cell(row=r, column=3, value='·伙伴：坐骑、翅膀、守护相关无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·助战：神像、契约、泰坦、奖杯、圣物无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·神兵：已激活的神兵转换职业后对应的神兵也会激活且继承升星星级，需确立对应神兵的一对一关系，器灵无损继承，器灵触发技能根据神兵替换对应的技能')
    r += 1
    ws.cell(row=r, column=3, value='·装备收藏：收藏的装备转换成对应职业的装备，再生属性无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·神佑：英雄圣装、神盾、神甲、神纹相关无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·伴侣：无损继承')
    r += 1
    ws.cell(row=r, column=3, value='·公会：保持不变')
    r += 1
    ws.cell(row=r, column=3, value='·交易行：交易行商品保持和之前相同的上架状态和上架时间，本职业上架的装备和道具转换为新职业对应的装备和道具')
    r += 1
    ws.cell(row=r, column=3, value='·商店：商店里面与本职业相关的道具转换为对应的新职业道具，购物单里面与本职业相关的道具转换为对应的新职业道具')
    r += 1
    ws.cell(row=r, column=3, value='·新职业外观：主城雕像、排行榜、天神骑士团、位面霸主、竞技场结算、王者1v1、结缘等显示需要匹配新职业')
    r += 1
    ws.cell(row=r, column=3, value='·奖励：福利、直购、运营活动、BP、寻宝、首充等于职业相关的需要匹配新职业')
    r += 1
    ws.cell(row=r, column=3, value='·转职：对应转职的职业称号等需要匹配以及继承对应的转职阶段，进行中的转职需要刷新为对应的新职业')
    r += 1
    ws.cell(row=r, column=3, value='·穿戴：保留玩家原本穿戴模型，包括装备、神兵、时装等外显')
    r += 1
    ws.cell(row=r, column=3, value='·成就：匹配对应转职的职业成就')
    r += 2
    
    ws.cell(row=r, column=2, value='界面预览')
    r += 1
    ws.cell(row=r, column=2, value='·放入给定的示意图，不要嵌入单元格，按照宽11cm的标准放入示意图，示意图不要重叠，每行最多展示4张图')
    r += 80  # 预留示意图空间
    
    # 功能入口
    ws.cell(row=r, column=3, value='功能入口')
    r += 2
    ws.cell(row=r, column=5, value='说明示例：')
    r += 1
    ws.cell(row=r, column=5, value='·入口位置：角色界面 → 新增页签"职业转换"，放在最后面')
    r += 1
    ws.cell(row=r, column=5, value='·背包里使用：点击跳转至职业转换界面')
    r += 1
    ws.cell(row=r, column=5, value='·功能开关开启时显示入口，关闭时隐藏入口')
    r += 4
    
    # 功能开关
    ws.cell(row=r, column=3, value='功能开关')
    r += 3
    ws.cell(row=r, column=3, value='·数值添加')
    r += 5
    
    # 功能详解
    ws.cell(row=r, column=2, value='·功能详解（下面是举例说明，左边大图为对应界面的示意图，右边为详细的功能点说明，说明时采用上面为该功能点的截图，下面为详细的说明）')
    r += 3
    
    # 职业转换界面详解
    ws.cell(row=r, column=15, value='·职业转换界面')
    r += 3
    ws.cell(row=r, column=15, value='·所有职业')
    r += 1
    ws.cell(row=r, column=16, value='·展示所有可转换的职业列表')
    r += 1
    ws.cell(row=r, column=16, value='·本职业置底，灰化不可选')
    r += 1
    ws.cell(row=r, column=16, value='·点击职业可查看该职业详情')
    r += 1
    ws.cell(row=r, column=16, value='·选中职业有高亮标识')
    r += 3
    
    ws.cell(row=r, column=15, value='·职业技能展示')
    r += 1
    ws.cell(row=r, column=16, value='·展示各职业所有主动技能')
    r += 1
    ws.cell(row=r, column=16, value='·技能简易描述')
    r += 1
    ws.cell(row=r, column=16, value='·各职业模型展示（用于展示技能效果）')
    r += 3
    
    ws.cell(row=r, column=15, value='·转职规则说明')
    r += 1
    ws.cell(row=r, column=16, value='·展示转职后的数据继承规则')
    r += 1
    ws.cell(row=r, column=16, value='·展示转职注意事项')
    r += 3
    
    ws.cell(row=r, column=15, value='·转职消耗')
    r += 1
    ws.cell(row=r, column=16, value='·显示转职所需消耗道具及数量')
    r += 1
    ws.cell(row=r, column=16, value='·道具不足时飘字提示')
    r += 3
    
    ws.cell(row=r, column=15, value='·转职按钮')
    r += 1
    ws.cell(row=r, column=16, value='·点击后判断所有前置条件是否满足')
    r += 1
    ws.cell(row=r, column=17, value='·条件满足：消耗道具，弹出二次确认弹窗')
    r += 1
    ws.cell(row=r, column=18, value='·确认后进入转职进度界面')
    r += 1
    ws.cell(row=r, column=18, value='·转职成功后弹窗提示"转职成功，需要重启游戏"')
    r += 1
    ws.cell(row=r, column=18, value='·关闭弹窗后重启游戏')
    r += 1
    ws.cell(row=r, column=17, value='·条件不满足：对应飘字提醒')
    r += 1
    ws.cell(row=r, column=16, value='·前置条件包括：可转换职业、可使用等级、开服天数限制、功能开关开启、非转职冷却时间、非战斗状态、非本职业、在主城、消耗道具足够')
    r += 3
    
    ws.cell(row=r, column=15, value='·转职进度界面')
    r += 1
    ws.cell(row=r, column=16, value='·展示转职进度动画')
    r += 1
    ws.cell(row=r, column=16, value='·期间玩家不能进行任何操作')
    r += 1
    ws.cell(row=r, column=16, value='·转职期间关掉游戏后，登录时如果转职未完成则继续弹转职进行中界面')
    r += 1
    ws.cell(row=r, column=16, value='·转职已完成则按照正常流程进入游戏')
    r += 5
    
    # 红点
    ws.cell(row=r, column=2, value='红点')
    r += 1
    ws.cell(row=r, column=2, value='·职业转换红点')
    r += 1
    ws.cell(row=r, column=3, value='·红点位置：角色界面页签、背包转换卡图标')
    r += 1
    ws.cell(row=r, column=3, value='·出现条件：拥有可使用的职业转换卡且满足使用条件')
    r += 1
    ws.cell(row=r, column=3, value='·消失条件：没有可使用的转换卡或功能开关关闭')
    r += 2
    
    # 公告
    ws.cell(row=r, column=2, value='公告')
    r += 1
    ws.cell(row=r, column=2, value='·转职成功后发送系统公告"[玩家名字]成功转职为[新职业名称]，实力大增！"')
    r += 2
    
    # 邮件
    ws.cell(row=r, column=2, value='邮件')
    r += 1
    ws.cell(row=r, column=2, value='·转职过程中数据异常时，通过邮件补发相关道具')
    r += 1
    ws.cell(row=r, column=3, value='·标题：职业转换补偿')
    r += 1
    ws.cell(row=r, column=3, value='·内容：您在职业转换过程中遇到异常，以下是为您补发的道具。')
    r += 2
    
    # 设置所有单元格字体
    for row in ws.iter_rows():
        for cell in row:
            cell.font = normal_font
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    
    # 保存
    output_path = r'E:\桌面\其它\职业转换卡系统策划案_模版格式.xlsx'
    wb.save(output_path)
    print(f'策划案已生成: {output_path}')

if __name__ == '__main__':
    create_doc()
