# -*- coding: utf-8 -*-
from openpyxl import load_workbook
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = load_workbook(r'E:\桌面\其它\策划案模版_完善版.xlsx')
ws = wb.active

# 打印模板内容结构
print('=== 模板结构 ===')
for row_idx in range(1, 90):
    row_content = []
    for col_idx in range(1, 10):
        val = ws.cell(row=row_idx, column=col_idx).value
        if val and str(val).strip():
            row_content.append('col{}={}'.format(col_idx, str(val)[:50]))
    if row_content:
        print('row{}: {}'.format(row_idx, ' | '.join(row_content)))
