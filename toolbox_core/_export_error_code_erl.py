#!/usr/bin/env python3
"""
独立脚本：读取 ErrorMessage.xlsm，生成 config/cfg_errorMessage.erl + .hrl

相当于 ExcelTool2.exe C++ ErlangService::writeFiles 的替代实现。
通过 subprocess 被 web_app.py 的 _exec_export_error_code 调用。
"""
import argparse
import os
import sys


HRL_TEMPLATE = '''-ifndef(cfg_errorMessage_hrl).
-define(cfg_errorMessage_hrl, true).

-record(errorMessageCfg, {{
        iD,
        errorString
}}).

-endif.
'''


def _read_xlsm_data(xlsm_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsm_path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for xr, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        a = str(row[0]).strip().lower() if row[0] is not None else ""
        if not a or a == "no":
            continue
        c = row[2]
        if c is None or not str(c).strip():
            continue
        b = row[1]
        if b is None:
            continue
        try:
            rid = int(float(str(b)))
        except Exception:
            continue
        rows.append((xr, rid, str(c)))
        if a == "end":
            break
    wb.close()
    return rows


def _write_erl(entries, config_dir):
    os.makedirs(config_dir, exist_ok=True)
    N = '\r\n'
    fp = os.path.join(config_dir, 'cfg_errorMessage.erl')
    with open(fp, 'w', encoding='utf-8', newline='') as f:
        ks = [e[1] for e in entries]
        f.write(f'-module(cfg_errorMessage).{N}')
        f.write(f'-include("cfg_errorMessage.hrl").{N}')
        f.write(f'-export([row/1, first_row/0, last_row/0, rows/1, rows/0, keys_length/0]).{N}')
        f.write(f'-export([getRow/1, getKeyList/0]).{N}')
        f.write(f'{N}%% 指定行{N}')
        f.write(f'row(Id) -> getRow(Id).{N}')
        f.write(f'{N}%% 第一行、最后一行{N}')
        f.write(f'first_row() -> getRow(0).{N}')
        if ks:
            f.write(f'last_row() -> getRow({ks[-1]}).{N}')
        f.write(f'{N}%% 行列表{N}')
        f.write(f'rows(KeyList) -> [row(Key) || Key <- KeyList].{N}')
        f.write(f'rows() -> rows(getKeyList()).{N}')
        f.write(f'{N}%% Key列表长度{N}')
        f.write(f'keys_length() -> {len(ks)}.{N}{N}')
        for _, rid, text in entries:
            escaped = text.replace('\\', '\\\\').replace('"', '\\"')
            f.write(f'getRow({rid}) -> #errorMessageCfg{{{N}')
            f.write(f'\tiD = {rid},{N}')
            f.write(f'\terrorString = "{escaped}"}};{N}')
        f.write(f'getRow(_) ->{N}\t{{}}.{N}{N}')
        f.write(f'getKeyList() -> [{N}')
        for i2, k in enumerate(ks):
            sep = f'].{N}' if i2 + 1 >= len(ks) else f',{N}'
            f.write(f'\t{k}{sep}')
    print(f"  [ERL] {fp}")

    hp = os.path.join(config_dir, 'cfg_errorMessage.hrl')
    with open(hp, 'w', encoding='utf-8', newline='') as f:
        for line in ['-ifndef(cfg_errorMessage_hrl).', '-define(cfg_errorMessage_hrl, true).', '',
                     '-record(errorMessageCfg, {', '\tiD,', '\terrorString', '}).', '', '-endif.']:
            f.write(line + '\r\n')
    print(f"  [HRL] {hp}")


def main():
    ap = argparse.ArgumentParser(description="ErrorMessage erlang 导出")
    ap.add_argument("--xlsm", required=True, help="ErrorMessage.xlsm 路径")
    ap.add_argument("--lang-dir", required=True, help="语言目录路径")
    args = ap.parse_args()

    xlsm_path = os.path.abspath(args.xlsm)
    lang_dir = os.path.abspath(args.lang_dir)
    config_dir = os.path.join(lang_dir, "config")

    if not os.path.isfile(xlsm_path):
        print(f"文件不存在: {xlsm_path}")
        sys.exit(1)

    try:
        entries = _read_xlsm_data(xlsm_path)
    except Exception as e:
        print(f"读取 xlsm 失败: {e}")
        sys.exit(1)

    if not entries:
        print("没有读取到有效数据")
        sys.exit(1)

    _write_erl(entries, config_dir)
    print(f"成功导出 {len(entries)} 条错误码 erlang")
    sys.exit(0)


if __name__ == "__main__":
    main()
