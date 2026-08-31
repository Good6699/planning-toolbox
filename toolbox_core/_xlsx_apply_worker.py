#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""openpyxl 合并子进程 worker：把大表 openpyxl 处理移到子进程，避免占主进程 GIL。
用法: _xlsx_apply_worker.py <arg_pickle> <res_pickle>
args 结构: {"mode": "merge_sheet_rows"|"copy_rows_by_id"|"sync_index", ...}
结果 pickle: {"ok": bool, "added": int, "updated": int, "logs": [str]}"""
import sys
import os
import pickle

_script_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _script_dir)


def main():
    if len(sys.argv) != 3:
        sys.stderr.write("Usage: _xlsx_apply_worker.py <arg_pickle> <res_pickle>\n")
        sys.exit(1)
    arg_path, res_path = sys.argv[1], sys.argv[2]
    try:
        with open(arg_path, "rb") as f:
            args = pickle.load(f)
    except Exception as e:
        sys.stderr.write(f"[worker] args 读取失败: {e}\n")
        sys.exit(1)

    import openpyxl
    from toolbox_xlsx_merge import _merge_sheet_rows, _copy_rows_by_id, _sync_index_table

    logs = []

    def put(msg):
        logs.append(msg)
        try:
            sys.stdout.write(msg)
            sys.stdout.flush()
        except Exception:
            pass

    result = {"ok": False, "added": 0, "updated": 0, "logs": logs}
    try:
        mode = args.get("mode")
        if mode == "merge_sheet_rows":
            # 语义合并应用阶段：差异数据 → 目标表（与 merge_texts_xlsm 原逻辑一致）
            target_path = args["target_path"]
            title_rows = args["title_rows"]
            id_col = args["id_col"]
            diff_data = args.get("diff_data") or {}   # {sheet: [row_dict, ...]}
            header_data = args.get("header_data") or {}  # {sheet: {行号: {列字母: 表头值}}}
            wb_tgt = openpyxl.load_workbook(target_path, keep_vba=True)
            total_added = total_updated = 0
            try:
                for sheet_name, sheet_rows in diff_data.items():
                    if sheet_name not in wb_tgt.sheetnames:
                        continue
                    ws_tgt = wb_tgt[sheet_name]
                    wb_tmp = openpyxl.Workbook()
                    ws_tmp = wb_tmp.active
                    ws_tmp.title = sheet_name
                    sheet_hd = header_data.get(sheet_name, {})
                    for row_num, cols in sheet_hd.items():
                        for letter, val in cols.items():
                            col_num = openpyxl.utils.column_index_from_string(letter)
                            ws_tmp.cell(row=row_num, column=col_num, value=val)
                    col_name_map = {}
                    if title_rows in sheet_hd:
                        for letter, hdr_val in sheet_hd[title_rows].items():
                            col_name_map[hdr_val] = openpyxl.utils.column_index_from_string(letter)
                    put(f"  {sheet_name}: {len(sheet_rows)} 行差异\n")
                    skip_keys = {"操作", "当前版本", "上一版本", "前一版本", "sheet",
                                 "_id_changed", "前一版本_ID", "前一版本_SC", "前一版本_sub"}
                    inp_rows = []
                    for i, row_data in enumerate(sheet_rows):
                        if row_data.get("操作") == "删除":
                            continue
                        r = title_rows + 1 + len(inp_rows)
                        inp_rows.append(r)
                        for col_name, val in row_data.items():
                            if col_name in skip_keys:
                                continue
                            col_num = col_name_map.get(col_name)
                            if col_num:
                                ws_tmp.cell(row=r, column=col_num, value=val)
                    if inp_rows:
                        a, u, _ = _merge_sheet_rows(ws_tmp, ws_tgt, inp_rows, title_rows, id_col, set(), put)
                        total_added += a
                        total_updated += u
                        put(f"  {sheet_name}: {u} 修改, {a} 新增\n")
                    wb_tmp.close()
            finally:
                # 仅在有实际修改时才保存，避免 openpyxl 无谓重写导致 svn 误标 M
                if total_added + total_updated > 0:
                    wb_tgt.save(target_path)
                wb_tgt.close()
            result.update(ok=True, added=total_added, updated=total_updated)
        elif mode == "copy_rows_by_id":
            # 指定合并文字表：按 ID 整行复制
            wb_src = openpyxl.load_workbook(args["src_path"], read_only=True, data_only=True)
            wb_tgt = openpyxl.load_workbook(args["tgt_path"], keep_vba=True)
            a = u = 0
            try:
                a, u = _copy_rows_by_id(wb_src, wb_tgt, args["id_by_sheet"],
                                        args["title_rows"], args["id_col"], put)
            finally:
                # 仅在有实际修改时才保存，避免 openpyxl 无谓重写导致 svn 误标 M
                if a + u > 0:
                    wb_tgt.save(args["tgt_path"])
                wb_src.close()
                wb_tgt.close()
            result.update(ok=True, added=a, updated=u)
        elif mode == "sync_index":
            a, u = _sync_index_table(args["src_idx"], args["tgt_idx"], args["changed_ids"], put)
            result.update(ok=True, added=a, updated=u)
        else:
            put(f"未知模式: {mode}\n")
    except Exception as e:
        import traceback
        put(f"worker 异常: {e}\n{traceback.format_exc()}\n")
        result["ok"] = False
    with open(res_path, "wb") as f:
        pickle.dump(result, f)
    sys.exit(0)


if __name__ == "__main__":
    main()
