#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
导出错误码为 Erlang 格式 (.erl + .hrl)
用法:
  python _export_error_code_erl.py --xlsm <xlsm路径> --lang-dir <语言目录>
输出:
  <lang-dir>/config/cfg_errorMessage.erl
  <lang-dir>/config/cfg_errorMessage.hrl
"""
import argparse
import os
import sys


def _read_xlsm_data(xlsm_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsm_path, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for xr, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        a = str(row[0]).strip().lower() if row[0] is not None else ""
        if not a or a == "no":
            continue
        c = row[2]
        if c is None or not str(c).strip():
            continue
        b = row[1]
        if b is None:
            continue
        try:
            rid = int(float(str(b)))
        except Exception:
            continue
        rows.append((xr, rid, str(c)))
        if a == "end":
            break
    wb.close()
    return rows


def _write_erl(entries, path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    keys = [e[1] for e in entries]
    with open(path, "w", encoding="utf-8", newline="\r\n") as f:
        f.write("-module(cfg_errorMessage).\r\n")
        f.write('-include("cfg_errorMessage.hrl").\r\n')
        f.write("-export([row/1, first_row/0, last_row/0, rows/1, rows/0, keys_length/0]).\r\n")
        f.write("-export([getRow/1, getKeyList/0]).\r\n")
        f.write("\r\n%% \\u6307\\u5b9a\\u884c\r\n")
        f.write("row(Id) -> getRow(Id).\r\n\r\n")
        f.write("%% \\u7b2c\\u4e00\\u884c\\u3001\\u6700\\u540e\\u4e00\\u884c\r\n")
        f.write("first_row() -> getRow(0).\r\n")
        if keys:
            f.write(f"last_row() -> getRow({keys[-1]}).\r\n")
        f.write("\r\n%% \\u884c\\u5217\\u8868\r\n")
        f.write("rows(KeyList) -> [row(Key) || Key <- KeyList].\r\n")
        f.write("rows() -> rows(getKeyList()).\r\n\r\n")
        f.write("%% Key\\u5217\\u8868\\u957f\\u5ea6\r\n")
        f.write(f"keys_length() -> {len(keys)}.\r\n\r\n")
        for entry in entries:
            rid, text = entry[1], entry[2]
            escaped = text.replace("\\", "\\\\").replace('"', '\\"')
            f.write(f"getRow({rid}) -> #errorMessageCfg{{\r\n")
            f.write(f"\tiD = {rid},\r\n")
            f.write(f'\terrorString = "{escaped}"}};\r\n')
        f.write("getKeyList() -> [\r\n")
        items = [str(k) for k in keys]
        for i in range(0, len(items), 20):
            chunk = items[i:i + 20]
            sep = "].\r\n" if i + 20 >= len(items) else ",\r\n"
            f.write("\t" + ",".join(chunk) + sep)


def main():
    parser = argparse.ArgumentParser(description="导出错误码为 Erlang 格式")
    parser.add_argument("--xlsm", required=True, help="ErrorMessage.xlsm 路径")
    parser.add_argument("--lang-dir", required=True, help="语言目录（含 config/）")
    args = parser.parse_args()

    entries = _read_xlsm_data(args.xlsm)
    if not entries:
        print("没有读取到有效数据")
        return 1

    hrl_path = os.path.join(args.lang_dir, "config", "cfg_errorMessage.hrl")
    os.makedirs(os.path.dirname(hrl_path), exist_ok=True)
    with open(hrl_path, "w", encoding="utf-8", newline="\r\n") as f:
        f.write("-ifndef(cfg_errorMessage_hrl).\r\n")
        f.write("-define(cfg_errorMessage_hrl, true).\r\n\r\n")
        f.write("-record(errorMessageCfg, {\r\n")
        f.write("\tiD,\r\n")
        f.write("\terrorString\r\n")
        f.write("}).\r\n\r\n")
        f.write("-endif.\r\n")

    erl_path = os.path.join(args.lang_dir, "config", "cfg_errorMessage.erl")
    _write_erl(entries, erl_path)

    print(f"已导出 {len(entries)} 条错误码")
    print(f"  {erl_path}")
    print(f"  {hrl_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
