#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""策划工具箱 - SVN工作流页签"""
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, font as tkfont
import json, os, subprocess, sys as _sys, threading, re, shutil, stat, copy, concurrent.futures
from datetime import datetime
from toolbox_platform import _DropTarget, _check_office_lock, _get_subprocess_kwargs, _get_svn_path
from toolbox_config import CONFIG_FILE, SCRIPT_DIR, MAIN_SCRIPT, DEFAULT_OUTPUT_DIR, load_config, save_config, int_or
from xlsm_zipper import apply_via_excel

class WorkflowTabMixin:
    def _wf_clear_detail(self):
        for w in self.wf_detail_frame.winfo_children():
            w.destroy()
        self.wf_detail_frame.columnconfigure(0, weight=1)

    def _wf_show_workflow_detail(self, wf_idx):
        self._wf_clear_detail()
        wf = self._wf_data[wf_idx]
        self.wf_detail_frame.columnconfigure(0, weight=1)

        tk.Label(self.wf_detail_frame, text="工作流设置",
                 font=("微软雅黑", 10, "bold")).pack(anchor="w", pady=(0, 8))

        name_frame = tk.Frame(self.wf_detail_frame)
        name_frame.pack(fill="x")
        name_frame.columnconfigure(1, weight=1)
        tk.Label(name_frame, text="名称：", font=("微软雅黑", 9)).grid(row=0, column=0, sticky="w")
        wf_name_var = tk.StringVar(value=wf.get("name", ""))
        def _on_name_change(*_):
            wf["name"] = wf_name_var.get()
            self._wf_save_config()
            if self.wf_tree.exists(f"wf_{wf_idx}"):
                self.wf_tree.item(f"wf_{wf_idx}", text=wf["name"])
        wf_name_var.trace_add("write", _on_name_change)
        ttk.Entry(name_frame, textvariable=wf_name_var, font=("微软雅黑", 9)
                  ).grid(row=0, column=1, sticky="ew", padx=(5, 0))

        ttk.Separator(self.wf_detail_frame, orient="horizontal").pack(fill="x", pady=8)

        tk.Label(self.wf_detail_frame, text="执行步骤",
                 font=("微软雅黑", 9, "bold")).pack(anchor="w")

        steps = wf.get("steps", [])
        if steps:
            step_list_frame = tk.Frame(self.wf_detail_frame)
            step_list_frame.pack(fill="x", pady=(4, 0))
            step_list_frame.columnconfigure(0, weight=1)
            step_list_frame.rowconfigure(0, weight=1)
            self.wf_step_listbox = tk.Listbox(step_list_frame, font=("微软雅黑", 9),
                                              height=12, activestyle="none")
            self.wf_step_listbox.grid(row=0, column=0, sticky="nsew")
            step_vbar = ttk.Scrollbar(step_list_frame, orient="vertical",
                                       command=self.wf_step_listbox.yview)
            step_vbar.grid(row=0, column=1, sticky="ns")
            self.wf_step_listbox.configure(yscrollcommand=step_vbar.set)
            for i, s in enumerate(steps):
                st = s.get("type", "")
                t = {"export_text": "导出文字表", "upload_svn": "上传SVN", "merge_table": "合并表格", "merge_translation": "合并翻译", "export_error_code": "导出错误码", "lock_svn": "锁定SVN", "open_tables": "打开表格"}.get(st, st)
                self.wf_step_listbox.insert("end", f"{i+1}. {t} - {s.get('name', '')}")
            sel_idx = self._selected_step if self._selected_step is not None else 0
            if sel_idx < len(steps):
                self.wf_step_listbox.selection_set(sel_idx)
            self.wf_step_listbox.bind("<Double-Button-1>", lambda e: self._wf_step_list_double_click(wf_idx))
        else:
            tk.Label(self.wf_detail_frame, text="(暂无步骤，点击下方添加)",
                     font=("微软雅黑", 9), fg="#888").pack(anchor="w", pady=(4, 0))

        btn_frame = tk.Frame(self.wf_detail_frame)
        btn_frame.pack(anchor="w", pady=(6, 0))

        self.wf_add_step_btn = ttk.Menubutton(btn_frame, text="+ 添加步骤", width=12)
        self.wf_add_step_btn.pack(side="left", padx=(0, 6))
        step_menu = tk.Menu(self.wf_add_step_btn, tearoff=0)
        step_menu.add_command(label="导出文字表", command=lambda: self._wf_add_step(wf_idx, "export_text"))
        step_menu.add_command(label="上传SVN", command=lambda: self._wf_add_step(wf_idx, "upload_svn"))
        step_menu.add_command(label="合并表格", command=lambda: self._wf_add_step(wf_idx, "merge_table"))
        step_menu.add_command(label="合并翻译", command=lambda: self._wf_add_step(wf_idx, "merge_translation"))
        step_menu.add_command(label="导出错误码", command=lambda: self._wf_add_step(wf_idx, "export_error_code"))
        step_menu.add_command(label="锁定SVN", command=lambda: self._wf_add_step(wf_idx, "lock_svn"))
        step_menu.add_command(label="打开表格", command=lambda: self._wf_add_step(wf_idx, "open_tables"))
        self.wf_add_step_btn.config(menu=step_menu)

        self.wf_del_step_btn = ttk.Button(btn_frame, text="- 删除选中步骤",
                                          command=lambda: self._wf_del_step(wf_idx), width=12)
        self.wf_del_step_btn.pack(side="left")

        ttk.Separator(self.wf_detail_frame, orient="horizontal").pack(fill="x", pady=8)

        exec_frame = tk.Frame(self.wf_detail_frame)
        exec_frame.pack(fill="x", pady=(10, 0))
        if steps:
            inner = tk.Frame(exec_frame)
            inner.pack(anchor="center")
            inner.columnconfigure(1, weight=0)

            tk.Label(inner, text="步骤范围：", font=("微软雅黑", 9)).grid(row=0, column=0, sticky="w")
            wf_step_range_var = tk.StringVar()
            wf_step_range_entry = ttk.Entry(inner, textvariable=wf_step_range_var,
                                            font=("微软雅黑", 9), width=14)
            wf_step_range_entry.grid(row=0, column=1, padx=(4, 12))
            wf_step_range_entry.insert(0, "")
            tk.Label(inner, text="如 1-3,5 或留空=全部",
                     font=("微软雅黑", 8), fg="#888").grid(row=0, column=2, padx=(0, 12))

            tk.Button(inner, text="执行工作流",
                      font=("微软雅黑", 10, "bold"),
                      bg="#2980b9", fg="white",
                      activebackground="#3498db",
                      relief="flat", padx=16, pady=4,
                      cursor="hand2",
                      command=lambda: threading.Thread(
                           target=self._wf_execute_workflow, args=(wf_idx, wf_step_range_var), daemon=True).start()).grid(row=0, column=3)

        self._selected_wf = wf_idx
        self._selected_step = None

    def _wf_show_step_detail(self, wf_idx, step_idx):
        self._wf_clear_detail()
        wf = self._wf_data[wf_idx]
        step = wf["steps"][step_idx]
        step_type = step.get("type", "export_text")
        self.wf_detail_frame.columnconfigure(0, weight=1)

        title_frame = tk.Frame(self.wf_detail_frame)
        title_frame.pack(fill="x")
        tk.Label(title_frame, text="步骤设置",
                 font=("微软雅黑", 10, "bold")).pack(side="left")
        ttk.Button(title_frame, text="<- 返回工作流",
                   command=lambda: self._wf_select_workflow_in_tree(wf_idx)).pack(side="right")

        grp1 = tk.LabelFrame(self.wf_detail_frame, text="  基本设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp1.pack(fill="x", pady=(8, 0))
        grp1.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp1, text="类型：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        type_map = {"export_text": "导出文字表", "upload_svn": "上传SVN", "merge_table": "合并表格", "merge_translation": "合并翻译", "export_error_code": "导出错误码", "lock_svn": "锁定SVN", "open_tables": "打开表格"}
        tk.Label(grp1, text=type_map.get(step_type, step_type),
                 font=("微软雅黑", 9), fg="#555").grid(row=row, column=1, sticky="w", padx=(5, 0))

        row = 1
        tk.Label(grp1, text="名称：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(4, 0))
        name_var = tk.StringVar(value=step.get("name", ""))
        def _save_name(*_):
            step["name"] = name_var.get()
            self._wf_save_config()
            self._wf_update_step_listbox(wf_idx)
            iid = "wf_" + str(wf_idx) + "_step_" + str(step_idx)
            if hasattr(self, "wf_tree") and self.wf_tree.exists(iid):
                st = step.get("type", "")
                t = {"export_text": "导出文字表", "upload_svn": "上传SVN", "merge_table": "合并表格", "merge_translation": "合并翻译", "export_error_code": "导出错误码", "lock_svn": "锁定SVN", "open_tables": "打开表格"}.get(st, st)
                self.wf_tree.item(iid, text=str(step_idx + 1) + ". " + t + " - " + step.get("name", ""))
        name_var.trace_add("write", lambda *_: _save_name())
        ttk.Entry(grp1, textvariable=name_var, font=("微软雅黑", 9)).grid(row=row, column=1, sticky="ew", padx=(5, 0), pady=(4, 0))

        if step_type == "export_text":
            self._wf_show_export_text_config(wf_idx, step_idx, step)
        elif step_type == "upload_svn":
            self._wf_show_upload_svn_config(wf_idx, step_idx, step)
        elif step_type == "merge_table":
            self._wf_show_merge_table_config(wf_idx, step_idx, step)
        elif step_type == "merge_translation":
            self._wf_show_merge_translation_config(wf_idx, step_idx, step)
        elif step_type == "export_error_code":
            self._wf_show_export_error_code_config(wf_idx, step_idx, step)
        elif step_type == "lock_svn":
            self._wf_show_lock_svn_config(wf_idx, step_idx, step)
        elif step_type == "open_tables":
            self._wf_show_open_tables_config(wf_idx, step_idx, step)

        exec_frame = tk.Frame(self.wf_detail_frame)
        exec_frame.pack(fill="x", pady=(10, 0))
        tk.Button(exec_frame, text="执行此步骤",
                  font=("微软雅黑", 10, "bold"),
                  bg="#27ae60", fg="white",
                  activebackground="#2ecc71",
                  relief="flat", padx=16, pady=4,
                  cursor="hand2",
                  command=lambda: threading.Thread(
                      target=self._wf_execute_single_step, args=(wf_idx, step_idx), daemon=True).start()).pack()

        self._selected_wf = wf_idx
        self._selected_step = step_idx

    def _wf_show_export_text_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  导出文字表设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp2, text="输入文件：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        input_var = tk.StringVar(value=step.get("input_file", ""))
        input_entry = ttk.Entry(grp2, textvariable=input_var, font=("微软雅黑", 9))
        input_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0))
        def _on_input_drop(files):
            if files:
                f = files[0].strip('"').strip("'")
                input_var.set(f)
                step["input_file"] = f
                self._wf_save_config()
        _DropTarget(input_entry, _on_input_drop).hook()
        def _browse_input():
            f = filedialog.askopenfilename(title="选择Excel文件",
                filetypes=[("Excel文件", "*.xlsm *.xlsx *.xls"), ("所有文件", "*.*")])
            if f:
                input_var.set(f)
                step["input_file"] = f
                self._wf_save_config()
        def _save_input(*_):
            step["input_file"] = input_var.get()
            self._wf_save_config()
        input_var.trace_add("write", lambda *_: _save_input())
        ttk.Button(grp2, text="浏览...", command=_browse_input, width=6
                   ).grid(row=row, column=2, padx=(5, 0))

        row = 1
        tk.Label(grp2, text="使用工具：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(4, 0))
        tools_frame = tk.Frame(grp2)
        tools_frame.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(5, 0), pady=(4, 0))
        tools_frame.columnconfigure(0, weight=1)

        tools_listbox = tk.Listbox(tools_frame, font=("Consolas", 9), height=3, activestyle="none")
        tools_listbox.grid(row=0, column=0, sticky="ew")
        for t in step.get("tools", []):
            tools_listbox.insert("end", os.path.basename(t) if t else t)

        def _on_tools_drop(files):
            for f in files:
                f = f.strip('"').strip("'")
                if f:
                    tools_listbox.insert("end", os.path.basename(f))
                    tools = step.setdefault("tools", [])
                    tools.append(f)
            self._wf_save_config()
            self._wf_update_step_listbox(wf_idx)
        _DropTarget(tools_listbox, _on_tools_drop).hook()

        def _add_tool():
            f = filedialog.askopenfilename(title="选择工具",
                filetypes=[("可执行文件", "*.exe *.bat *.cmd"), ("所有文件", "*.*")])
            if f:
                tools_listbox.insert("end", os.path.basename(f))
                tools = step.setdefault("tools", [])
                tools.append(f)
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)

        def _remove_tool():
            sel = tools_listbox.curselection()
            if sel:
                idx = sel[0]
                tools_listbox.delete(idx)
                tools = step.get("tools", [])
                if idx < len(tools):
                    tools.pop(idx)
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        btn2 = tk.Frame(tools_frame)
        btn2.grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Button(btn2, text="添加", command=_add_tool, width=5).pack(side="left", padx=(0, 5))
        ttk.Button(btn2, text="移除", command=_remove_tool, width=5).pack(side="left")

    def _wf_show_upload_svn_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  上传SVN设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp2, text="上传目录：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        dirs_frame = tk.Frame(grp2)
        dirs_frame.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(5, 0))
        dirs_frame.columnconfigure(0, weight=1)

        dirs_listbox = tk.Listbox(dirs_frame, font=("Consolas", 9), height=3, activestyle="none")
        dirs_listbox.grid(row=0, column=0, sticky="ew")
        for d in step.get("dirs", []):
            dirs_listbox.insert("end", d)

        def _on_dirs_drop(files):
            for f in files:
                f = f.strip('"').strip("'")
                if os.path.isdir(f):
                    dirs_listbox.insert("end", f)
                    dirs = step.setdefault("dirs", [])
                    if f not in dirs:
                        dirs.append(f)
            self._wf_save_config()
            self._wf_update_step_listbox(wf_idx)
        _DropTarget(dirs_listbox, _on_dirs_drop).hook()

        def _add_dir():
            d = filedialog.askdirectory(title="选择上传目录")
            if d:
                dirs_listbox.insert("end", d)
                dirs = step.setdefault("dirs", [])
                dirs.append(d)
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)

        def _remove_dir():
            sel = dirs_listbox.curselection()
            if sel:
                idx = sel[0]
                dirs_listbox.delete(idx)
                dirs = step.get("dirs", [])
                if idx < len(dirs):
                    dirs.pop(idx)
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        btn2 = tk.Frame(dirs_frame)
        btn2.grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Button(btn2, text="添加", command=_add_dir, width=5).pack(side="left", padx=(0, 5))
        ttk.Button(btn2, text="移除", command=_remove_dir, width=5).pack(side="left")

    def _wf_show_merge_table_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  合并表格设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        # ── 输入1：文件/目录列表 ──
        row = 0
        tk.Label(grp2, text="输入文件：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        paths_frame = tk.Frame(grp2)
        paths_frame.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(5, 0))
        paths_frame.columnconfigure(0, weight=1)

        paths_listbox = tk.Listbox(paths_frame, font=("Consolas", 9), height=3, activestyle="none")
        paths_listbox.grid(row=0, column=0, sticky="ew")

        import tkinter.font as tkfont
        _listbox_font = tkfont.Font(font=paths_listbox.cget("font"))

        def _get_available_chars():
            w = paths_listbox.winfo_width()
            if w < 50:
                return 70
            cw = _listbox_font.measure("0")
            if cw < 1:
                return 70
            return max(25, (w // cw) - 3)

        def _shorten_path(p):
            max_chars = _get_available_chars()
            if len(p) <= max_chars:
                return p
            parts = p.split(os.sep)
            if len(parts) <= 3:
                return p
            head = os.sep.join(parts[:3])
            tail = parts[-1]
            ellipsis = "..."
            needed_for_tail = len(ellipsis + os.sep + tail)
            head_max = max_chars - needed_for_tail
            if head_max < 3:
                head_max = 3
            head = head[:head_max]
            return head + os.sep + ellipsis + os.sep + tail

        def _refresh_listbox_display():
            paths = step.get("input_paths", [])
            paths_listbox.delete(0, "end")
            for p in paths:
                paths_listbox.insert("end", _shorten_path(p))
        paths_listbox.bind("<Configure>", lambda e: _refresh_listbox_display())
        _refresh_listbox_display()

        def _normalize_path(p):
            p = p.strip('"').strip("'")
            return os.path.normpath(p)

        def _on_paths_drop(files):
            existing = step.setdefault("input_paths", [])
            existing_norm = {os.path.normpath(x) for x in existing}
            changed = False
            for f in files:
                f = _normalize_path(f)
                if f and f not in existing_norm:
                    existing.append(f)
                    existing_norm.add(f)
                    changed = True
            if changed:
                _refresh_listbox_display()
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)
        _DropTarget(paths_listbox, _on_paths_drop).hook()

        def _get_default_dir():
            paths = step.get("input_paths", [])
            if paths:
                last = os.path.normpath(paths[-1])
                if os.path.isdir(last):
                    return last
                else:
                    d = os.path.dirname(last)
                    if os.path.isdir(d):
                        return d
            return None

        def _add_path():
            initial = _get_default_dir()
            f = filedialog.askopenfilename(title="选择Excel文件",
                initialdir=initial,
                filetypes=[("Excel文件", "*.xlsm *.xlsx *.xls"), ("所有文件", "*.*")])
            if f:
                f = os.path.normpath(f)
                existing = step.setdefault("input_paths", [])
                existing_norm = {os.path.normpath(x) for x in existing}
                if f not in existing_norm:
                    existing.append(f)
                    _refresh_listbox_display()
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        def _add_dir_path():
            initial = _get_default_dir()
            d = filedialog.askdirectory(title="选择目录（将扫描其中所有Excel文件）",
                initialdir=initial)
            if not d:
                return
            existing = step.setdefault("input_paths", [])
            existing_norm = {os.path.normpath(x) for x in existing}
            found = 0
            for root, _dirs, files in os.walk(d):
                for f in files:
                    ext = os.path.splitext(f)[1].lower()
                    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
                        full = os.path.normpath(os.path.join(root, f))
                        if full not in existing_norm:
                            existing.append(full)
                            existing_norm.add(full)
                            found += 1
            if found == 0:
                d = os.path.normpath(d)
                if d not in existing_norm:
                    existing.append(d)
            _refresh_listbox_display()
            self._wf_save_config()
            self._wf_update_step_listbox(wf_idx)

        def _remove_path():
            sel = paths_listbox.curselection()
            if sel:
                idx = sel[0]
                paths = step.get("input_paths", [])
                if idx < len(paths):
                    paths.pop(idx)
                    _refresh_listbox_display()
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        def _clear_all_paths():
            if not step.get("input_paths", []):
                return
            if messagebox.askyesno("确认清空", "确定要删除所有输入文件吗？"):
                step["input_paths"] = []
                _refresh_listbox_display()
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)

        btn1 = tk.Frame(paths_frame)
        btn1.grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Button(btn1, text="添加文件", command=_add_path, width=8).pack(side="left", padx=(0, 3))
        ttk.Button(btn1, text="添加目录", command=_add_dir_path, width=8).pack(side="left", padx=(0, 3))
        ttk.Button(btn1, text="移除", command=_remove_path, width=5).pack(side="left", padx=(0, 3))
        ttk.Button(btn1, text="清空全部", command=_clear_all_paths, width=8).pack(side="left")

        # ── 输入2：目标文件夹 ──
        row = 1
        tk.Label(grp2, text="目标文件夹：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        tgt_var = tk.StringVar(value=step.get("target_dir", ""))
        tgt_entry = ttk.Entry(grp2, textvariable=tgt_var, font=("微软雅黑", 9))
        tgt_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0), pady=(6, 0))

        def _on_tgt_drop(files):
            if files:
                f = files[0].strip('"').strip("'")
                if os.path.isdir(f):
                    tgt_var.set(f)
                    step["target_dir"] = f
                    self._wf_save_config()
        _DropTarget(tgt_entry, _on_tgt_drop).hook()

        def _browse_tgt():
            d = filedialog.askdirectory(title="选择目标文件夹")
            if d:
                tgt_var.set(d)
                step["target_dir"] = d
                self._wf_save_config()

        def _save_tgt(*_):
            step["target_dir"] = tgt_var.get()
            self._wf_save_config()
        tgt_var.trace_add("write", lambda *_: _save_tgt())
        ttk.Button(grp2, text="浏览...", command=_browse_tgt, width=6
                   ).grid(row=row, column=2, padx=(5, 0), pady=(6, 0))

        # ── 过滤前缀配置 ──
        row = 2
        tk.Label(grp2, text="过滤前缀：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        prefix_frame = tk.Frame(grp2)
        prefix_frame.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(5, 0), pady=(6, 0))
        prefix_frame.columnconfigure(0, weight=1)

        prefix_listbox = tk.Listbox(prefix_frame, font=("Consolas", 9), height=3, activestyle="none")
        prefix_listbox.grid(row=0, column=0, sticky="ew")
        for p in step.get("merge_prefixes", []):
            prefix_listbox.insert("end", p)

        prefix_var = tk.StringVar()
        prefix_entry = ttk.Entry(prefix_frame, textvariable=prefix_var, font=("微软雅黑", 9))
        prefix_entry.grid(row=1, column=0, sticky="ew", pady=(4, 0))

        def _add_prefix():
            val = prefix_var.get().strip()
            if val and val not in step.get("merge_prefixes", []):
                prefix_listbox.insert("end", val)
                prefixes = step.setdefault("merge_prefixes", [])
                prefixes.append(val)
                prefix_var.set("")
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)

        def _remove_prefix():
            sel = prefix_listbox.curselection()
            if sel:
                idx = sel[0]
                prefix_listbox.delete(idx)
                prefixes = step.get("merge_prefixes", [])
                if idx < len(prefixes):
                    prefixes.pop(idx)
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        prefix_entry.bind("<Return>", lambda e: _add_prefix())
        btn_p = tk.Frame(prefix_frame)
        btn_p.grid(row=2, column=0, sticky="w", pady=(4, 0))
        ttk.Button(btn_p, text="添加", command=_add_prefix, width=5).pack(side="left", padx=(0, 3))
        ttk.Button(btn_p, text="移除", command=_remove_prefix, width=5).pack(side="left")

        tk.Label(grp2, text="匹配时将去掉前缀再进行精准匹配（区分大小写）",
                 font=("微软雅黑", 8), fg="#888").grid(row=3, column=1, sticky="w", padx=(5, 0))

        # ── 标题行数 + ID列（独立配置） ──
        row = 4
        tk.Label(grp2, text="标题行数：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        title_rows_var = tk.StringVar(value=str(step.get("title_rows", "1")))
        title_rows_entry = ttk.Entry(grp2, textvariable=title_rows_var, font=("微软雅黑", 9), width=10)
        title_rows_entry.grid(row=row, column=1, sticky="w", padx=(5, 0), pady=(6, 0))
        def _save_title_rows(*_):
            step["title_rows"] = title_rows_var.get().strip()
            self._wf_save_config()
        title_rows_var.trace_add("write", lambda *_: _save_title_rows())

        row = 5
        tk.Label(grp2, text="ID 列号：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        id_col_var = tk.StringVar(value=str(step.get("id_col", "1")))
        id_col_spin = ttk.Spinbox(grp2, from_=1, to=100, textvariable=id_col_var, width=6, font=("微软雅黑", 9))
        id_col_spin.grid(row=row, column=1, sticky="w", padx=(5, 0), pady=(6, 0))
        def _save_id_col(*_):
            step["id_col"] = id_col_var.get().strip()
            self._wf_save_config()
        id_col_var.trace_add("write", lambda *_: _save_id_col())
        tk.Label(grp2, text="标题行与该列交叉的单元格的值作为 ID",
                 font=("微软雅黑", 8), fg="#888").grid(row=row, column=1, columnspan=2, sticky="w", padx=(90, 0), pady=(6, 0))

    def _wf_show_export_error_code_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  导出错误码设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp2, text="根目录：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        root_var = tk.StringVar(value=step.get("root_dir", ""))
        root_entry = ttk.Entry(grp2, textvariable=root_var, font=("微软雅黑", 9))
        root_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0))
        def _on_root_drop(files):
            if files:
                f = files[0].strip('"').strip("'")
                root_var.set(f)
                step["root_dir"] = f
                self._wf_save_config()
        _DropTarget(root_entry, _on_root_drop).hook()
        def _browse_root():
            d = filedialog.askdirectory(title="选择根目录（自动查找 gameData\\Language）")
            if d:
                root_var.set(d)
                step["root_dir"] = d
                self._wf_save_config()
        def _save_root(*_):
            step["root_dir"] = root_var.get()
            self._wf_save_config()
        root_var.trace_add("write", lambda *_: _save_root())
        ttk.Button(grp2, text="浏览...", command=_browse_root, width=6
                   ).grid(row=row, column=2, padx=(5, 0))

        row = 1
        tk.Label(grp2, text="项目根目录，自动查找 gameData\\Language",
                 font=("微软雅黑", 8), fg="#888").grid(row=row, column=0, columnspan=3, sticky="w", padx=(5, 0))

        row = 2
        tk.Label(grp2, text="语言列表：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        lang_var = tk.StringVar(value=step.get("lang_codes", ""))
        lang_entry = ttk.Entry(grp2, textvariable=lang_var, font=("微软雅黑", 9))
        lang_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0), pady=(6, 0))
        def _save_lang(*_):
            step["lang_codes"] = lang_var.get()
            self._wf_save_config()
        lang_var.trace_add("write", lambda *_: _save_lang())

        row = 3
        tk.Label(grp2, text="逗号分隔，如 JA_JP,KO_KR,ZH_CN,EN_US",
                 font=("微软雅黑", 8), fg="#888").grid(row=row, column=0, columnspan=3, sticky="w", padx=(5, 0))

    def _wf_show_lock_svn_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  锁定SVN设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp2, text="目标路径：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        target_var = tk.StringVar(value=step.get("target_path", ""))
        target_entry = ttk.Entry(grp2, textvariable=target_var, font=("微软雅黑", 9))
        target_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0))
        def _browse_target():
            f = filedialog.askopenfilename(title="选择要锁定的SVN文件")
            if not f:
                f = filedialog.askdirectory(title="选择要锁定的SVN文件夹")
            if f:
                target_var.set(f)
                step["target_path"] = f
                self._wf_save_config()
        ttk.Button(grp2, text="浏览…", command=_browse_target,
                   width=8).grid(row=row, column=2, padx=(5, 0))
        def _on_target_drop(files):
            if files:
                f = files[0].strip('"').strip("'")
                target_var.set(f)
                step["target_path"] = f
                self._wf_save_config()
        _DropTarget(target_entry, _on_target_drop).hook()
        def _save_target(*_):
            step["target_path"] = target_var.get()
            self._wf_save_config()
        target_var.trace_add("write", lambda *_: _save_target())

        row = 1
        tk.Label(grp2, text="锁定备注：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        msg_var = tk.StringVar(value=step.get("lock_msg", "锁定中，请勿修改"))
        msg_entry = ttk.Entry(grp2, textvariable=msg_var, font=("微软雅黑", 9))
        msg_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0), pady=(6, 0))
        def _save_msg(*_):
            step["lock_msg"] = msg_var.get()
            self._wf_save_config()
        msg_var.trace_add("write", lambda *_: _save_msg())

        row = 2
        tk.Label(grp2, text="更新目录：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        dirs_frame = tk.Frame(grp2)
        dirs_frame.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(5, 0), pady=(6, 0))
        dirs_frame.columnconfigure(0, weight=1)

        dirs_listbox = tk.Listbox(dirs_frame, font=("Consolas", 9), height=3, activestyle="none")
        dirs_listbox.grid(row=0, column=0, sticky="ew")
        for d in step.get("update_dirs", []):
            dirs_listbox.insert("end", d)

        def _on_dirs_drop(files):
            for f in files:
                f = f.strip('"').strip("'")
                if os.path.isdir(f):
                    dirs_listbox.insert("end", f)
                    dirs = step.setdefault("update_dirs", [])
                    if f not in dirs:
                        dirs.append(f)
            self._wf_save_config()
        _DropTarget(dirs_listbox, _on_dirs_drop).hook()

        def _add_dir():
            d = filedialog.askdirectory(title="选择要更新的SVN目录")
            if d:
                dirs_listbox.insert("end", d)
                dirs = step.setdefault("update_dirs", [])
                dirs.append(d)
                self._wf_save_config()

        def _remove_dir():
            sel = dirs_listbox.curselection()
            if sel:
                idx = sel[0]
                dirs_listbox.delete(idx)
                dirs = step.get("update_dirs", [])
                if idx < len(dirs):
                    dirs.pop(idx)
                    self._wf_save_config()

        btn2 = tk.Frame(dirs_frame)
        btn2.grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Button(btn2, text="添加", command=_add_dir, width=5).pack(side="left", padx=(0, 5))
        ttk.Button(btn2, text="移除", command=_remove_dir, width=5).pack(side="left")

        row = 3
        tip_text = "选中后执行 svn lock 操作，锁定该文件或文件夹。指定更新目录则锁前先 svn update --accept theirs-full。"
        tk.Label(grp2, text=tip_text,
                 font=("微软雅黑", 8), fg="#888", justify="left").grid(
            row=row, column=0, columnspan=3, sticky="w", padx=(5, 0), pady=(6, 0))

    def _wf_show_open_tables_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  打开表格设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp2, text="文件列表：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        paths_frame = tk.Frame(grp2)
        paths_frame.grid(row=row, column=1, columnspan=2, sticky="ew", padx=(5, 0))
        paths_frame.columnconfigure(0, weight=1)

        paths_listbox = tk.Listbox(paths_frame, font=("Consolas", 9), height=5, activestyle="none")
        paths_listbox.grid(row=0, column=0, sticky="ew")
        for p in step.get("file_paths", []):
            paths_listbox.insert("end", p)

        def _on_paths_drop(files):
            existing = step.setdefault("file_paths", [])
            for f in files:
                f = f.strip('"').strip("'")
                if f and f not in existing:
                    paths_listbox.insert("end", f)
                    existing.append(f)
            self._wf_save_config()
            self._wf_update_step_listbox(wf_idx)
        _DropTarget(paths_listbox, _on_paths_drop).hook()

        def _add_file():
            f = filedialog.askopenfilename(title="选择Excel文件",
                filetypes=[("Excel文件", "*.xlsm *.xlsx *.xls *.csv"), ("所有文件", "*.*")])
            if f:
                existing = step.setdefault("file_paths", [])
                if f not in existing:
                    paths_listbox.insert("end", f)
                    existing.append(f)
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        def _add_dir():
            d = filedialog.askdirectory(title="选择目录（将扫描其中所有Excel文件）")
            if not d:
                return
            existing = step.setdefault("file_paths", [])
            found = 0
            for root, _dirs, files in os.walk(d):
                for f in files:
                    ext = os.path.splitext(f)[1].lower()
                    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
                        full = os.path.normpath(os.path.join(root, f))
                        if full not in existing:
                            paths_listbox.insert("end", full)
                            existing.append(full)
                            found += 1
            if found:
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)

        def _remove_selected():
            sel = paths_listbox.curselection()
            if sel:
                idx = sel[0]
                paths_listbox.delete(idx)
                existing = step.get("file_paths", [])
                if idx < len(existing):
                    existing.pop(idx)
                    self._wf_save_config()
                    self._wf_update_step_listbox(wf_idx)

        def _clear_all():
            if not step.get("file_paths", []):
                return
            if messagebox.askyesno("确认清空", "确定要删除所有文件吗？"):
                step["file_paths"] = []
                paths_listbox.delete(0, "end")
                self._wf_save_config()
                self._wf_update_step_listbox(wf_idx)

        btn1 = tk.Frame(paths_frame)
        btn1.grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Button(btn1, text="添加文件", command=_add_file, width=8).pack(side="left", padx=(0, 3))
        ttk.Button(btn1, text="添加目录", command=_add_dir, width=8).pack(side="left", padx=(0, 3))
        ttk.Button(btn1, text="移除", command=_remove_selected, width=5).pack(side="left", padx=(0, 3))
        ttk.Button(btn1, text="清空全部", command=_clear_all, width=8).pack(side="left")

        tk.Label(grp2, text="支持拖拽文件添加，执行时用系统默认程序批量打开所有文件",
                 font=("微软雅黑", 8), fg="#888").grid(row=2, column=0, columnspan=3, sticky="w", padx=(5, 0), pady=(6, 0))

    def _wf_show_merge_translation_config(self, wf_idx, step_idx, step):
        grp2 = tk.LabelFrame(self.wf_detail_frame, text="  合并翻译设置  ",
                              font=("微软雅黑", 9), padx=8, pady=6)
        grp2.pack(fill="x", pady=(8, 0))
        grp2.columnconfigure(1, weight=1)

        row = 0
        tk.Label(grp2, text="翻译文件：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w")
        input_var = tk.StringVar(value=step.get("input_file", ""))
        input_entry = ttk.Entry(grp2, textvariable=input_var, font=("微软雅黑", 9))
        input_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0))
        def _on_input_drop(files):
            if files:
                f = files[0].strip('"').strip("'")
                input_var.set(f)
                step["input_file"] = f
                self._wf_save_config()
        _DropTarget(input_entry, _on_input_drop).hook()
        def _browse_input():
            f = filedialog.askopenfilename(title="选择翻译Excel文件",
                filetypes=[("Excel文件", "*.xlsm *.xlsx *.xls"), ("所有文件", "*.*")])
            if f:
                input_var.set(f)
                step["input_file"] = f
                self._wf_save_config()
        def _save_input(*_):
            step["input_file"] = input_var.get()
            self._wf_save_config()
        input_var.trace_add("write", lambda *_: _save_input())
        ttk.Button(grp2, text="浏览...", command=_browse_input, width=6
                   ).grid(row=row, column=2, padx=(5, 0))

        row = 1
        tk.Label(grp2, text="原文件：", font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        orig_var = tk.StringVar(value=step.get("original_file", ""))
        orig_entry = ttk.Entry(grp2, textvariable=orig_var, font=("微软雅黑", 9))
        orig_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0), pady=(6, 0))
        def _on_orig_drop(files):
            if files:
                f = files[0].strip('"').strip("'")
                orig_var.set(f)
                step["original_file"] = f
                self._wf_save_config()
        _DropTarget(orig_entry, _on_orig_drop).hook()
        def _browse_orig():
            f = filedialog.askopenfilename(title="选择原文件",
                filetypes=[("Excel文件", "*.xlsm *.xlsx *.xls"), ("所有文件", "*.*")])
            if f:
                orig_var.set(f)
                step["original_file"] = f
                self._wf_save_config()
        def _save_orig(*_):
            step["original_file"] = orig_var.get()
            self._wf_save_config()
        orig_var.trace_add("write", lambda *_: _save_orig())
        ttk.Button(grp2, text="浏览...", command=_browse_orig, width=6
                   ).grid(row=row, column=2, padx=(5, 0), pady=(6, 0))

        row = 2
        tk.Label(grp2, text="Sheet名（留空=所有共有Sheet）：",
                 font=("微软雅黑", 9)).grid(row=row, column=0, sticky="w", pady=(6, 0))
        sheet_var = tk.StringVar(value=step.get("sheet_name", ""))
        sheet_entry = ttk.Entry(grp2, textvariable=sheet_var, font=("微软雅黑", 9))
        sheet_entry.grid(row=row, column=1, sticky="ew", padx=(5, 0), pady=(6, 0))
        def _save_sheet(*_):
            step["sheet_name"] = sheet_var.get()
            self._wf_save_config()
        sheet_var.trace_add("write", lambda *_: _save_sheet())

        row = 3
        tk.Label(grp2, text="按ID匹配行，按表头名匹配列，将翻译内容写入原文件对应单元格",
                 font=("微软雅黑", 8), fg="#888").grid(row=row, column=0, columnspan=3, sticky="w", padx=(5, 0))

    def _wf_update_step_listbox(self, wf_idx, select_idx=None):
        if hasattr(self, "wf_step_listbox") and self.wf_step_listbox.winfo_exists():
            self.wf_step_listbox.delete(0, "end")
            wf = self._wf_data[wf_idx]
            for i, s in enumerate(wf.get("steps", [])):
                st = s.get("type", "")
                t = {"export_text": "导出文字表", "upload_svn": "上传SVN", "merge_table": "合并表格", "merge_translation": "合并翻译", "export_error_code": "导出错误码", "lock_svn": "锁定SVN", "open_tables": "打开表格"}.get(st, st)
                self.wf_step_listbox.insert("end", f"{i+1}. {t} - {s.get('name', '')}")
            if select_idx is not None and select_idx < self.wf_step_listbox.size():
                self.wf_step_listbox.selection_set(select_idx)

    def _wf_step_list_double_click(self, wf_idx):
        sel = self.wf_step_listbox.curselection()
        if sel:
            self._wf_select_step_in_tree(wf_idx, sel[0])
            self._on_wf_tree_select()

    def _wf_add_step(self, wf_idx, step_type):
        wf = self._wf_data[wf_idx]
        steps = wf.setdefault("steps", [])
        if len(steps) >= 20:
            self._wlog("最多20个步骤", "warn")
            return
        n = len(steps) + 1
        if step_type == "export_text":
            steps.append({"type": "export_text", "name": str(n), "input_file": "", "tools": []})
        elif step_type == "upload_svn":
            steps.append({"type": "upload_svn", "name": str(n), "dirs": []})
        elif step_type == "export_error_code":
            steps.append({"type": "export_error_code", "name": str(n), "root_dir": "", "lang_codes": ""})
        elif step_type == "lock_svn":
            steps.append({"type": "lock_svn", "name": str(n), "target_path": "", "update_dirs": [], "lock_msg": "锁定中，请勿修改"})
        elif step_type == "merge_translation":
            steps.append({"type": "merge_translation", "name": str(n), "input_file": "", "original_file": "", "sheet_name": ""})
        elif step_type == "open_tables":
            steps.append({"type": "open_tables", "name": str(n), "file_paths": []})
        else:
            steps.append({"type": "merge_table", "name": str(n), "input_paths": [], "target_dir": "", "merge_prefixes": [], "title_rows": "1", "id_col": "1"})
        self._wf_refresh_tree(select_iid="wf_" + str(wf_idx) + "_step_" + str(len(steps) - 1))
        self._wf_save_config()
        self._wlog("已添加步骤: " + steps[-1]["name"], "ok")

    def _wf_del_step(self, wf_idx):
        if not hasattr(self, "wf_step_listbox") or not self.wf_step_listbox.winfo_exists():
            return
        sel = self.wf_step_listbox.curselection()
        if not sel:
            return
        step_idx = sel[0]
        wf = self._wf_data[wf_idx]
        steps = wf.get("steps", [])
        if step_idx < len(steps):
            name = steps[step_idx].get("name", "")
            if messagebox.askyesno("确认删除", "删除步骤[" + name + "]?"):
                del steps[step_idx]
                self._wf_refresh_tree(select_iid="wf_" + str(wf_idx))
                self._wf_save_config()
                self._wlog("已删除步骤: " + name, "warn")

    def _wf_move_step(self, wf_idx, direction):
        if not hasattr(self, "wf_step_listbox") or not self.wf_step_listbox.winfo_exists():
            return
        sel = self.wf_step_listbox.curselection()
        if not sel:
            return
        step_idx = sel[0]
        target = step_idx + direction
        if target < 0 or target >= len(self._wf_data[wf_idx].get("steps", [])):
            return
        steps = self._wf_data[wf_idx]["steps"]
        steps[step_idx], steps[target] = steps[target], steps[step_idx]
        self._wf_save_config()
        self._wf_update_step_listbox(wf_idx)
        self._wf_refresh_tree(select_iid="wf_" + str(wf_idx) + "_step_" + str(target))
        self.wf_step_listbox.selection_set(target)
        self._wlog("步骤已移动", "info")

    def _wf_load_workflows(self):
        self._wf_data = list(self.config.get("workflows", []))
        self._selected_wf = None
        self._selected_step = None
        self._wf_refresh_tree()

    def _wf_save_config(self):
        self.config["workflows"] = self._wf_data
        save_config(self.config)
        self.config = load_config()

    def _wf_refresh_tree(self, select_iid=None):
        sel = self.wf_tree.selection()
        saved = select_iid if select_iid else (sel[0] if sel else None)
        for item in self.wf_tree.get_children():
            self.wf_tree.delete(item)
        for i, wf in enumerate(self._wf_data):
            name = wf.get("name", "工作流" + str(i+1))
            pid = self.wf_tree.insert("", "end", text=name, iid="wf_" + str(i))
            for j, step in enumerate(wf.get("steps", [])):
                st = step.get("type", "")
                t = {"export_text": "导出文字表", "upload_svn": "上传SVN", "merge_table": "合并表格", "merge_translation": "合并翻译", "export_error_code": "导出错误码", "lock_svn": "锁定SVN", "open_tables": "打开表格"}.get(st, st)
                self.wf_tree.insert(pid, "end", text=str(j+1) + ". " + t + " - " + step.get("name", ""),
                                    iid="wf_" + str(i) + "_step_" + str(j))
        f = tkfont.Font(font=ttk.Style().lookup("Treeview", "font"))
        max_w = 0
        for item in self.wf_tree.get_children():
            w = f.measure(self.wf_tree.item(item, "text"))
            max_w = max(max_w, w)
            for child in self.wf_tree.get_children(item):
                w = f.measure(self.wf_tree.item(child, "text"))
                max_w = max(max_w, w + 20)
        self.wf_tree.column("#0", width=max_w + 24, stretch=False)
        if saved and self.wf_tree.exists(saved):
            if select_iid:
                if "_step_" in saved:
                    parts = saved.split("_")
                    parent_iid = "wf_" + parts[1]
                    if self.wf_tree.exists(parent_iid):
                        self.wf_tree.item(parent_iid, open=True)
                else:
                    self.wf_tree.item(saved, open=True)
            self.wf_tree.selection_set(saved)
            self.wf_tree.focus(saved)
            self.root.after_idle(self._on_wf_tree_select)

    def _wf_select_workflow_in_tree(self, wf_idx):
        iid = "wf_" + str(wf_idx)
        self.wf_tree.selection_set(iid)
        self.wf_tree.focus(iid)
        self._on_wf_tree_select()

    def _wf_select_step_in_tree(self, wf_idx, step_idx):
        iid = "wf_" + str(wf_idx) + "_step_" + str(step_idx)
        self.wf_tree.selection_set(iid)
        self.wf_tree.focus(iid)

    def _wf_tree_drag_start(self, event):
        item = self.wf_tree.identify_row(event.y)
        if not item:
            self._wf_drag_item = None
            return
        if "_step_" in item:
            parent = self.wf_tree.parent(item)
            if not parent:
                self._wf_drag_item = None
                return
            self._wf_drag_item = item
            self._wf_drag_parent = parent
        else:
            parent = self.wf_tree.parent(item)
            if parent != "":
                self._wf_drag_item = None
                return
            self._wf_drag_item = item
            self._wf_drag_parent = None
        self._wf_drag_start_y = event.y

    def _wf_tree_drag_motion(self, event):
        if not self._wf_drag_item:
            return
        if abs(event.y - self._wf_drag_start_y) > 10:
            self.wf_tree.config(cursor="hand2")

    def _wf_tree_drag_end(self, event):
        if not self._wf_drag_item:
            return
        self.wf_tree.config(cursor="")
        if abs(event.y - self._wf_drag_start_y) < 15:
            self._wf_drag_item = None
            self._wf_drag_parent = None
            return

        target = self.wf_tree.identify_row(event.y)
        if not target:
            self._wf_drag_item = None
            self._wf_drag_parent = None
            return

        src_iid = self._wf_drag_item

        # ── 拖拽步骤（子项）──
        if "_step_" in src_iid:
            src_parts = src_iid.split("_")
            wf_idx = int(src_parts[1])
            src_step_idx = int(src_parts[3])

            if "_step_" in target:
                tgt_parts = target.split("_")
                tgt_wf_idx = int(tgt_parts[1])
                if tgt_wf_idx != wf_idx:
                    self._wf_drag_item = None
                    self._wf_drag_parent = None
                    return
                tgt_step_idx = int(tgt_parts[3])
            else:
                tgt_wf_idx = int(target.split("_")[1])
                if tgt_wf_idx != wf_idx:
                    self._wf_drag_item = None
                    self._wf_drag_parent = None
                    return
                tgt_step_idx = len(self._wf_data[wf_idx].get("steps", []))

            if src_step_idx == tgt_step_idx:
                self._wf_drag_item = None
                self._wf_drag_parent = None
                return

            steps = self._wf_data[wf_idx]["steps"]
            step_item = steps.pop(src_step_idx)
            if src_step_idx < tgt_step_idx:
                tgt_step_idx -= 1
            steps.insert(tgt_step_idx, step_item)
            self._wf_save_config()
            parent_iid = "wf_" + str(wf_idx)
            self._wf_refresh_tree(select_iid=parent_iid)
            self._wf_drag_item = None
            self._wf_drag_parent = None
            self._wlog("步骤顺序已调整", "info")
            return

        # ── 拖拽工作流（顶层）──
        src_idx = int(src_iid.split("_")[1])

        if "_step_" not in target:
            tgt_idx = int(target.split("_")[1])
        else:
            parts = target.split("_")
            tgt_idx = int(parts[1])

        if src_idx == tgt_idx:
            self._wf_drag_item = None
            self._wf_drag_parent = None
            return

        item = self._wf_data.pop(src_idx)
        if src_idx < tgt_idx:
            tgt_idx -= 1
        self._wf_data.insert(tgt_idx, item)
        self._wf_save_config()
        self._wf_refresh_tree(select_iid="wf_" + str(tgt_idx))
        self._wf_drag_item = None
        self._wf_drag_parent = None

    def _wf_show_create_dialog(self):
        dialog = tk.Toplevel(self.root)
        dialog.title("新建工作流")
        dialog.resizable(False, False)
        dialog.transient(self.root)
        dialog.grab_set()

        tk.Label(dialog, text="工作流名称：", font=("微软雅黑", 9)).grid(row=0, column=0, padx=15, pady=(15, 5), sticky="w")
        name_var = tk.StringVar()
        name_entry = ttk.Entry(dialog, textvariable=name_var, font=("微软雅黑", 9), width=30)
        name_entry.grid(row=1, column=0, padx=15, pady=(0, 15))
        name_entry.select_range(0, "end")
        name_entry.focus_set()

        btn_frame = tk.Frame(dialog)
        btn_frame.grid(row=2, column=0, pady=(0, 12))

        def _on_ok():
            name = name_var.get().strip()
            if not name:
                messagebox.showwarning("提示", "请输入工作流名称")
                return
            self._wf_data.append({"name": name, "steps": []})
            self._wf_refresh_tree()
            self._wf_save_config()
            self._wlog("已创建工作流: " + name, "ok")
            dialog.destroy()

        ttk.Button(btn_frame, text="确定", command=_on_ok, width=8).pack(side="left", padx=6)
        ttk.Button(btn_frame, text="取消", command=dialog.destroy, width=8).pack(side="left", padx=6)
        dialog.bind("<Return>", lambda e: _on_ok())
        x = self.root.winfo_rootx() + (self.root.winfo_width() - 280) // 2
        y = self.root.winfo_rooty() + (self.root.winfo_height() - 150) // 2
        dialog.geometry("+" + str(x) + "+" + str(y))

    def _wf_del_workflow(self):
        sel = self.wf_tree.selection()
        if not sel:
            return
        iid = sel[0]
        if iid.startswith("wf_") and "_step_" not in iid:
            idx = int(iid.split("_")[1])
            name = self._wf_data[idx].get("name", "")
            if messagebox.askyesno("确认删除", "删除工作流[" + name + "]及其所有步骤？"):
                del self._wf_data[idx]
                self._wf_refresh_tree()
                self._wf_save_config()
                self._wf_clear_detail()
                tk.Label(self.wf_detail_frame, text="请从左侧选择一个工作流",
                         font=("微软雅黑", 9), fg="#888").pack(pady=20)
                self._wlog("已删除工作流: " + name, "warn")
        elif "_step_" in iid:
            parts = iid.split("_")
            wf_idx = int(parts[1])
            step_idx = int(parts[3])
            wf = self._wf_data[wf_idx]
            steps = wf.get("steps", [])
            if step_idx < len(steps):
                name = steps[step_idx].get("name", "")
                if messagebox.askyesno("确认删除", "删除步骤[" + name + "]？"):
                    del steps[step_idx]
                    self._wf_refresh_tree(select_iid="wf_" + str(wf_idx))
                    self._wf_save_config()
                    self._wlog("已删除步骤: " + name, "warn")

    def _wf_copy_workflow(self):
        sel = self.wf_tree.selection()
        if not sel:
            return
        iid = sel[0]
        if iid.startswith("wf_") and "_step_" not in iid:
            idx = int(iid.split("_")[1])
            orig = self._wf_data[idx]
            new_wf = copy.deepcopy(orig)
            new_wf["name"] = orig.get("name", "") + "(副本)"
            insert_pos = idx + 1
            self._wf_data.insert(insert_pos, new_wf)
            self._wf_save_config()
            new_iid = "wf_" + str(insert_pos)
            self._wf_refresh_tree(select_iid=new_iid)
            self._wlog("已复制工作流: " + orig.get("name", "") + " -> " + new_wf["name"], "ok")

    def _wf_build_log_area(self):
        parent = self.workflow_tab
        parent.rowconfigure(1, weight=0)
        log_frame = tk.Frame(parent)
        log_frame.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(4, 0))
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(1, weight=1)

        log_header = tk.Frame(log_frame)
        log_header.grid(row=0, column=0, sticky="ew")
        log_header.columnconfigure(0, weight=1)
        tk.Label(log_header, text="执行日志",
                 font=("微软雅黑", 9, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Button(log_header, text="清空", command=self._wf_clear_log, width=5
                   ).grid(row=0, column=1, padx=(8, 0))

        self.wf_log_text = tk.Text(log_frame, height=8,
                                   font=("Consolas", 9),
                                   bg="#1e1e1e", fg="#d4d4d4",
                                   insertbackground="white",
                                   state="disabled", wrap="word")
        self.wf_log_text.grid(row=1, column=0, sticky="nsew", pady=(4, 0))
        log_scroll = ttk.Scrollbar(log_frame, orient="vertical",
                                    command=self.wf_log_text.yview)
        log_scroll.grid(row=1, column=1, sticky="ns", pady=(4, 0))
        self.wf_log_text["yscrollcommand"] = log_scroll.set

        for tag in ("info", "ok", "warn", "error", "head"):
            self.wf_log_text.tag_config(tag,
                foreground={"info": "#d4d4d4", "ok": "#4ec9b0", "warn": "#dcdcaa",
                            "error": "#f44747", "head": "#569cd6"}[tag])

    def _wf_clear_log(self):
        self.wf_log_text.config(state="normal")
        self.wf_log_text.delete("1.0", "end")
        self.wf_log_text.config(state="disabled")

    def _wlog(self, msg, level="info"):
        if not hasattr(self, "wf_log_text") or not self.wf_log_text.winfo_exists():
            return
        ts = datetime.now().strftime("%H:%M:%S")
        msg = f"[{ts}] {msg}"
        w = self.wf_log_text
        def _upd():
            try:
                w.config(state="normal")
                w.insert("end", msg + "\n", level)
                w.see("end")
                w.config(state="disabled")
            except Exception:
                pass
        self.root.after(0, _upd)

    def _wf_execute_workflow(self, wf_idx, step_range_var=None):
        wf = self._wf_data[wf_idx]
        steps = wf.get("steps", [])
        if not steps:
            self._wlog("没有步骤需要执行", "warn")
            return
        self._wlog("=" * 50, "head")
        self._wlog("执行工作流", "head")

        range_str = step_range_var.get().strip() if step_range_var else ""
        step_indices = self._parse_step_range(range_str, len(steps))
        if step_indices is None:
            self._wlog("步骤范围格式无效，请使用如 1-3,5 的格式", "error")
            return
        if len(step_indices) < len(steps):
            self._wlog("指定步骤: " + ",".join(str(i+1) for i in step_indices), "info")
        steps = [steps[i] for i in step_indices]

        blocked = False
        for i, step in enumerate(steps):
            self._wlog("-- [" + str(i+1) + "/" + str(len(steps)) + "] " + step.get("name", "") + " --", "head")
            if blocked:
                self._wlog("  已阻断，跳过", "warn")
                continue
            step_type = step.get("type", "")
            ok = True
            if step_type == "export_text":
                ok = self._wf_execute_export_text(step)
            elif step_type == "upload_svn":
                ok = self._wf_execute_upload_svn(step)
            elif step_type == "merge_table":
                ok = self._wf_execute_merge_table(step)
            elif step_type == "merge_translation":
                ok = self._wf_execute_merge_translation(step)
            elif step_type == "export_error_code":
                ok = self._wf_execute_export_error_code(step)
            elif step_type == "lock_svn":
                ok = self._wf_execute_lock_svn(step)
            elif step_type == "open_tables":
                ok = self._wf_execute_open_tables(step)
            else:
                self._wlog("未知步骤类型: " + str(step_type), "error")
                ok = False
            if not ok:
                self._wlog("❌ 步骤执行失败，阻断后续步骤", "error")
                blocked = True
        self._wlog("=" * 50, "head")
        if blocked:
            self._wf_clear_caches()
            self._wlog("缓存已清除", "ok")
        self._wlog("工作流执行完成", "ok")

    def _wf_execute_single_step(self, wf_idx, step_idx):
        wf = self._wf_data[wf_idx]
        step = wf["steps"][step_idx]
        self._wlog("=" * 50, "head")
        self._wlog("步骤: " + step.get("name", ""), "head")
        step_type = step.get("type", "")
        if step_type == "export_text":
            self._wf_execute_export_text(step)
        elif step_type == "upload_svn":
            self._wf_execute_upload_svn(step)
        elif step_type == "merge_table":
            self._wf_execute_merge_table(step)
        elif step_type == "merge_translation":
            self._wf_execute_merge_translation(step)
        elif step_type == "export_error_code":
            self._wf_execute_export_error_code(step)
        elif step_type == "lock_svn":
            self._wf_execute_lock_svn(step)
        elif step_type == "open_tables":
            self._wf_execute_open_tables(step)
        self._wlog("步骤执行完成", "ok")

    def _wf_execute_lock_svn(self, step):
        target = step.get("target_path", "").strip()
        lock_msg = step.get("lock_msg", "锁定中，请勿修改").strip()

        if not target or not os.path.exists(target):
            self._wlog("目标路径无效: " + str(target), "error")
            return False

        self._wlog("目标路径: " + target, "info")
        self._wlog("锁定备注: " + lock_msg, "info")

        try:
            svn_exe = _get_svn_path()

            if os.path.isfile(target):
                check_dir = os.path.dirname(target)
            else:
                check_dir = target

            info_result = subprocess.run(
                [svn_exe, "info", check_dir],
                capture_output=True, text=True,
                **(_get_subprocess_kwargs() if _sys.platform == "win32" else {})
            )
            if info_result.returncode != 0:
                err = info_result.stderr.strip()
                self._wlog("不是有效的 SVN 工作副本: " + err, "error")
                return False

            url_result = subprocess.run(
                [svn_exe, "info", "--show-item", "url", target],
                capture_output=True, text=True,
                **(_get_subprocess_kwargs() if _sys.platform == "win32" else {})
            )
            if url_result.returncode == 0:
                svn_url = url_result.stdout.strip()
                self._wlog("SVN URL: " + svn_url, "info")

            update_dirs = step.get("update_dirs", [])
            if update_dirs:
                for d in update_dirs:
                    d = d.strip()
                    if not d or not os.path.isdir(d):
                        self._wlog("更新目录无效: " + str(d), "error")
                        return False
                    self._wlog("正在更新目录: " + d, "info")
                    update_result = subprocess.run(
                        [svn_exe, "update", "--accept", "theirs-full", d],
                        capture_output=True, text=True,
                        **(_get_subprocess_kwargs() if _sys.platform == "win32" else {})
                    )
                    if update_result.returncode == 0:
                        out = update_result.stdout.strip()
                        self._wlog("✅ 更新完成: " + d, "ok")
                        for line in out.splitlines():
                            line = line.strip()
                            if line:
                                self._wlog("  " + line, "info")
                    else:
                        err = update_result.stderr.strip()
                        self._wlog("❌ 更新失败: " + d + " - " + err, "error")
                        return False

            lock_result = subprocess.run(
                [svn_exe, "lock", target, "-m", lock_msg],
                capture_output=True, text=True,
                **(_get_subprocess_kwargs() if _sys.platform == "win32" else {})
            )
            if lock_result.returncode == 0:
                self._wlog("✅ 锁定成功: " + target, "ok")
            else:
                err = lock_result.stderr.strip()
                self._wlog("❌ 锁定失败: " + err, "error")
                return False
        except Exception as e:
            self._wlog("锁定操作异常: " + str(e), "error")
            import traceback
            self._wlog(traceback.format_exc(), "error")
            return False
        return True

    def _wf_execute_export_text(self, step):
        tools = step.get("tools", [])
        if not tools:
            self._wlog("未配置工具，跳过", "warn")
            return True

        input_file = step.get("input_file", "").strip()
        if input_file:
            self._wlog("输入文件: " + os.path.basename(input_file), "info")
        self._wlog("使用 " + str(len(tools)) + " 个工具并行执行...", "info")

        def _run_one(tool_path):
            tool_path = tool_path.strip()
            if not tool_path or not os.path.exists(tool_path):
                self._wlog("工具不存在: " + str(tool_path), "warn")
                return
            try:
                ext = os.path.splitext(tool_path)[1].lower()
                if ext in (".bat", ".cmd"):
                    proc = subprocess.Popen(
                        ["cmd.exe", "/c", tool_path],
                        cwd=os.path.dirname(tool_path),
                        stdin=subprocess.PIPE,
                        creationflags=subprocess.CREATE_NEW_CONSOLE)
                else:
                    proc = subprocess.Popen(
                        [tool_path],
                        stdin=subprocess.PIPE,
                        creationflags=subprocess.CREATE_NEW_CONSOLE)
                proc.communicate(input=b"\n", timeout=3600)
                self._wlog("  " + os.path.basename(tool_path) + " 已完成", "ok")
            except subprocess.TimeoutExpired:
                try:
                    proc.kill()
                    proc.communicate(timeout=5)
                except:
                    pass
                self._wlog("  " + os.path.basename(tool_path) + " 超时", "error")
            except Exception as e:
                self._wlog("  " + os.path.basename(tool_path) + " 失败: " + str(e), "error")

        with concurrent.futures.ThreadPoolExecutor(max_workers=len(tools)) as executor:
            futures = [executor.submit(_run_one, t) for t in tools]
            concurrent.futures.wait(futures)

        self._wf_clear_caches()
        self._wlog("缓存已清除", "ok")
        return True

    def _wf_execute_open_tables(self, step):
        file_paths = step.get("file_paths", [])
        if not file_paths:
            self._wlog("没有要打开的文件", "warn")
            return True
        self._wlog("正在打开 " + str(len(file_paths)) + " 个文件...", "info")
        ok_count = 0
        for fp in file_paths:
            fp = fp.strip()
            if not fp or not os.path.exists(fp):
                self._wlog("文件不存在: " + fp, "warn")
                continue
            try:
                os.startfile(fp)
                self._wlog("已打开: " + os.path.basename(fp), "ok")
                ok_count += 1
            except Exception as e:
                self._wlog("打开失败: " + fp + " - " + str(e), "error")
        self._wlog("已打开 " + str(ok_count) + "/" + str(len(file_paths)) + " 个文件", "ok")
        return ok_count == len(file_paths)

    def _wf_execute_export_error_code(self, step):
        root_dir = step.get("root_dir", "").strip()
        lang_codes = step.get("lang_codes", "").strip()
        if not root_dir:
            self._wlog("未指定根目录", "error")
            return False
        if not lang_codes:
            self._wlog("未指定语言列表", "error")
            return False

        def _resolve_language_dir(base):
            base = os.path.abspath(base)
            if base.endswith("Language") and os.path.isdir(base):
                return base
            test = os.path.join(base, "Language")
            if os.path.isdir(test):
                return test
            test = os.path.join(base, "gameData", "Language")
            if os.path.isdir(test):
                return test
            return None

        lang_dir = _resolve_language_dir(root_dir)
        if not lang_dir:
            self._wlog("无法找到 gameData\\Language 目录: " + root_dir, "error")
            return False

        codes = [c.strip() for c in lang_codes.split(",") if c.strip()]
        if not codes:
            self._wlog("语言列表为空", "error")
            return False

        self._wlog("Language 目录: " + lang_dir, "info")
        self._wlog("处理语言: " + ", ".join(codes), "info")

        import subprocess as _sp

        script_dir = os.path.dirname(os.path.abspath(__file__))
        et2_path = os.path.join(script_dir, "ExcelTool2.py")
        et2_python = _sys.executable

        results = []

        for code in codes:
            lang_path = os.path.join(lang_dir, code)
            xlsm_file = os.path.join(lang_path, "Data2", "ErrorMessage.xlsm")

            if not os.path.isfile(xlsm_file):
                results.append((code, False, "ErrorMessage.xlsm 未找到"))
                self._wlog(f"  [{code}] SKIP: xlsm 未找到", "warn")
                continue

            # ExcelTool2.py 兼容两版源代码逻辑 (Asia 新版 4 参数 / KR2 旧版 2 参数)
            cmd = [et2_python, et2_path, "ErrorMessage", "--lang-dir", lang_path]
            try:
                r = _sp.run(cmd, capture_output=True, text=True, timeout=120)
                if r.returncode == 0:
                    self._wlog(f"  [{code}] {r.stdout.strip().split(chr(10))[-1]}", "ok")
                    results.append((code, True, "导出成功"))
                else:
                    err = (r.stderr or r.stdout or "").strip()[:200]
                    self._wlog(f"  [{code}] FAIL: {err}", "warn")
                    results.append((code, False, err))
            except Exception as e:
                self._wlog(f"  [{code}] ERROR: {e}", "error")
                results.append((code, False, str(e)))

        ok_count = sum(1 for _, ok, _ in results if ok)
        self._wlog(f"导出错误码完成: {ok_count}/{len(codes)}", "ok")
        if ok_count == len(codes):
            self._wlog("全部语言导出成功，proto/out + ExportTxt + Client 已更新", "ok")
            self._wlog("注意: config/*.erl 和 .ExcelTool2 由原 ExcelTool2.exe 管理，本步骤不覆盖", "info")

        self._wf_clear_caches()
        return ok_count == len(codes)

    def _wf_execute_merge_translation(self, step):
        excel_file = step.get("input_file", "").strip()
        original_file = step.get("original_file", "").strip()
        sheet_name = step.get("sheet_name", "").strip()

        if not excel_file or not os.path.exists(excel_file):
            self._wlog("翻译文件无效: " + str(excel_file), "error")
            return False
        if not original_file or not os.path.exists(original_file):
            self._wlog("原文件无效: " + str(original_file), "error")
            return False

        self._wlog("翻译文件: " + excel_file, "info")
        self._wlog("原文件: " + original_file, "info")

        import openpyxl

        _cfg = load_config()
        _lang_id_map = _cfg.get("tr_lang_id_map", {})
        _chinese_ids = set()
        for _ln, _ids in _lang_id_map.items():
            if "中文" in _ln or _ln.strip().lower() in ("chinese", "简体中文", "中文"):
                for _id in _ids:
                    _chinese_ids.add(_id.strip().lower())
        if not _chinese_ids:
            _chinese_ids = {"zh", "zh_cn", "zh-cn", "zh cn", "chinese", "简体中文", "中文", "简中", "cn"}

        try:
            trans_wb = openpyxl.load_workbook(excel_file, data_only=True)
            orig_wb = openpyxl.load_workbook(original_file)

            id_keys = ["id", "i_d", "编号", "key"]

            if sheet_name:
                if sheet_name not in trans_wb.sheetnames:
                    self._wlog("翻译文件中无 Sheet: " + sheet_name, "error")
                    trans_wb.close(); orig_wb.close()
                    return False
                if sheet_name not in orig_wb.sheetnames:
                    self._wlog("原文件中无 Sheet: " + sheet_name, "error")
                    trans_wb.close(); orig_wb.close()
                    return False
                process_sheets = [sheet_name]
            else:
                process_sheets = [sn for sn in trans_wb.sheetnames if sn in orig_wb.sheetnames]

            global_mode = False
            if not process_sheets:
                self._wlog("两文件无共有 Sheet，启用跨 Sheet ID 匹配模式", "info")
                global_mode = True

                orig_headers_by_sheet = {}
                orig_global_id_map = {}

                def _find_header_row(ws):
                    for r in range(1, min(ws.max_row, 6) + 1):
                        count = 0
                        for c in range(1, min(ws.max_column, 20) + 1):
                            v = ws.cell(row=r, column=c).value
                            if v is not None and isinstance(v, str) and v.strip():
                                count += 1
                        if count >= 3:
                            return r
                    return 1

                def _clean_header(name):
                    return name.strip().lower().strip(":").strip()

                for orig_sn in orig_wb.sheetnames:
                    orig_ws = orig_wb[orig_sn]
                    hdr_row = _find_header_row(orig_ws)
                    oh = {}
                    for col in range(1, orig_ws.max_column + 1):
                        h = orig_ws.cell(row=hdr_row, column=col).value
                        if h is not None:
                            cleaned = _clean_header(str(h))
                            if cleaned:
                                oh[cleaned] = col
                    orig_headers_by_sheet[orig_sn] = oh

                    orig_id_col_gs = None
                    for k in id_keys:
                        if k in oh:
                            orig_id_col_gs = oh[k]
                            break
                    if orig_id_col_gs is None:
                        continue

                    for row in range(hdr_row + 1, orig_ws.max_row + 1):
                        val = orig_ws.cell(row=row, column=orig_id_col_gs).value
                        if val is not None:
                            key = str(val).strip()
                            if key and key not in orig_global_id_map:
                                orig_global_id_map[key] = (orig_sn, row)

                process_sheets = trans_wb.sheetnames
                self._wlog("  扫描原文件 " + str(len(orig_headers_by_sheet)) + " 个 Sheet, " + str(len(orig_global_id_map)) + " 个 ID", "info")

            sheet_ops = []

            for sn in process_sheets:
                trans_ws = trans_wb[sn]

                trans_headers = {}
                for col in range(1, trans_ws.max_column + 1):
                    h = trans_ws.cell(row=1, column=col).value
                    if h is not None:
                        cleaned = str(h).strip().lower().strip(":").strip()
                        if cleaned:
                            trans_headers[cleaned] = col

                trans_id_col = None
                for k in id_keys:
                    if k in trans_headers:
                        trans_id_col = trans_headers[k]
                        break
                if trans_id_col is None and trans_headers:
                    first_key = next(iter(trans_headers))
                    trans_id_col = trans_headers[first_key]

                if global_mode:
                    header_col_map = {}
                    for h_name, t_col in trans_headers.items():
                        if t_col == trans_id_col:
                            continue
                        if h_name in _chinese_ids:
                            continue
                        for o_sn, oh in orig_headers_by_sheet.items():
                            if h_name in oh:
                                header_col_map.setdefault(o_sn, {})[t_col] = oh[h_name]
                    if not header_col_map:
                        self._wlog("  Sheet " + sn + ": 无匹配的翻译列", "warn")
                        continue

                    updates_by_orig_sheet = {}
                    matched_rows = 0
                    for row in range(2, trans_ws.max_row + 1):
                        id_val = trans_ws.cell(row=row, column=trans_id_col).value
                        if id_val is None:
                            continue
                        id_key = str(id_val).strip()
                        if not id_key or id_key not in orig_global_id_map:
                            continue
                        matched_rows += 1
                        orig_sn, orig_row = orig_global_id_map[id_key]
                        if orig_sn not in header_col_map:
                            continue
                        col_map = header_col_map[orig_sn]
                        orig_ws = orig_wb[orig_sn]
                        for t_col, o_col in col_map.items():
                            val = trans_ws.cell(row=row, column=t_col).value
                            if val is not None:
                                val_str = str(val)
                                orig_current = orig_ws.cell(row=orig_row, column=o_col).value
                                orig_current_str = str(orig_current or "")
                                if val_str.strip() != orig_current_str.strip():
                                    updates_by_orig_sheet.setdefault(orig_sn, []).append((orig_row, o_col, val))

                    for o_sn, upds in updates_by_orig_sheet.items():
                        sheet_ops.append({"sheet": o_sn, "updates": upds, "inserts": []})
                    total_upd = sum(len(v) for v in updates_by_orig_sheet.values())
                    if total_upd > 0:
                        self._wlog("  Sheet " + sn + ": 匹配 " + str(matched_rows) + " 行, " + str(total_upd) + " 个单元格待更新（跨 " + str(len(updates_by_orig_sheet)) + " 个原Sheet）", "info")
                    else:
                        self._wlog("  Sheet " + sn + ": 匹配 " + str(matched_rows) + " 行, 无变更", "info")
                else:
                    orig_ws = orig_wb[sn]

                    orig_headers = {}
                    for col in range(1, orig_ws.max_column + 1):
                        h = orig_ws.cell(row=1, column=col).value
                        if h is not None:
                            cleaned = str(h).strip().lower().strip(":").strip()
                            if cleaned:
                                orig_headers[cleaned] = col

                    orig_id_col = None
                    for k in id_keys:
                        if k in trans_headers and k in orig_headers:
                            orig_id_col = orig_headers[k]
                            break
                    if orig_id_col is None:
                        common = [k for k in trans_headers if k in orig_headers]
                        if common:
                            orig_id_col = orig_headers[common[0]]
                        else:
                            self._wlog("  Sheet " + sn + ": 无共有表头列，跳过", "warn")
                            continue

                    orig_id_map = {}
                    for row in range(2, orig_ws.max_row + 1):
                        val = orig_ws.cell(row=row, column=orig_id_col).value
                        if val is not None:
                            key = str(val).strip()
                            if key:
                                orig_id_map[key] = row

                    header_col_map = {}
                    for h_name, t_col in trans_headers.items():
                        if t_col == trans_id_col:
                            continue
                        if h_name in _chinese_ids:
                            continue
                        if h_name in orig_headers:
                            header_col_map[t_col] = orig_headers[h_name]

                    if not header_col_map:
                        self._wlog("  Sheet " + sn + ": 无匹配的翻译列", "warn")
                        continue

                    updates = []
                    matched_rows = 0
                    for row in range(2, trans_ws.max_row + 1):
                        id_val = trans_ws.cell(row=row, column=trans_id_col).value
                        if id_val is None:
                            continue
                        id_key = str(id_val).strip()
                        if not id_key or id_key not in orig_id_map:
                            continue
                        matched_rows += 1
                        orig_row = orig_id_map[id_key]
                        for t_col, o_col in header_col_map.items():
                            val = trans_ws.cell(row=row, column=t_col).value
                            if val is not None:
                                val_str = str(val)
                                orig_current = orig_ws.cell(row=orig_row, column=o_col).value
                                orig_current_str = str(orig_current or "")
                                if val_str.strip() != orig_current_str.strip():
                                    updates.append((orig_row, o_col, val))

                    if updates:
                        sheet_ops.append({"sheet": sn, "updates": updates, "inserts": []})
                        self._wlog("  Sheet " + sn + ": 匹配 " + str(matched_rows) + " 行, " + str(len(updates)) + " 个单元格待更新", "info")
                    else:
                        self._wlog("  Sheet " + sn + ": 匹配 " + str(matched_rows) + " 行, 无变更", "info")

            trans_wb.close()
            orig_wb.close()

            if not sheet_ops:
                self._wlog("没有需要更新的内容", "warn")
                return True

            if _check_office_lock(original_file):
                self._wlog("原文件被 WPS/Excel 锁定，无法保存", "error")
                return False

            self._wlog("Excel 后台写入中...", "info")
            ok, err_msg = apply_via_excel(original_file, sheet_ops)
            if ok:
                total_updates = sum(len(ops["updates"]) for ops in sheet_ops)
                self._wlog("已保存: " + os.path.basename(original_file), "ok")
                self._wlog("合并翻译完成: 更新 " + str(total_updates) + " 个单元格", "ok")
                return True
            else:
                self._wlog("Excel 写入失败: " + err_msg, "error")
                return False

        except Exception as e:
            self._wlog("处理失败: " + str(e), "error")
            import traceback
            self._wlog(traceback.format_exc(), "error")
            return False

    def _wf_clear_caches(self):
        cache_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cache")
        if os.path.isdir(cache_dir):
            import shutil as _su
            try:
                for f in os.listdir(cache_dir):
                    fp = os.path.join(cache_dir, f)
                    try:
                        if os.path.isfile(fp) or os.path.islink(fp):
                            os.remove(fp)
                        elif os.path.isdir(fp):
                            _su.rmtree(fp)
                    except Exception:
                        pass
            except Exception:
                pass
        try:
            import shutil as _su
            for p in ("__pycache__",):
                d = os.path.join(os.path.dirname(os.path.abspath(__file__)), p)
                if os.path.isdir(d):
                    _su.rmtree(d)
        except Exception:
            pass

    def _wf_execute_upload_svn(self, step):
        dirs = step.get("dirs", [])
        if not dirs:
            self._wlog("未配置上传目录，跳过", "warn")
            return True

        svn_exe = _get_svn_path()
        tortoise = self._get_tortoise_proc_path()
        if not tortoise:
            self._wlog("未找到 TortoiseSVN，无法提交。请安装 TortoiseSVN 后重试", "error")
            return False

        for d in dirs:
            d = d.strip()
            if not os.path.exists(d):
                self._wlog("路径不存在: " + d, "warn")
                continue
            if os.path.isfile(d):
                d = os.path.dirname(d)
            self._wlog("检查目录: " + d, "info")

            if svn_exe:
                self._wlog("正在更新目录: " + d, "info")
                _svn_enc = "gbk" if _sys.platform == "win32" else "utf-8"
                try:
                    r = subprocess.run(
                        [svn_exe, "update", "--accept", "theirs-full", d],
                        capture_output=True, text=True,
                        encoding=_svn_enc, errors="replace",
                        timeout=120, **_get_subprocess_kwargs()
                    )
                    if r.returncode == 0:
                        for line in r.stdout.strip().splitlines():
                            line = line.strip()
                            if line:
                                self._wlog("  " + line, "info")
                        self._wlog("✅ 更新完成: " + d, "ok")
                    else:
                        err = r.stderr.strip()
                        self._wlog("❌ 更新失败: " + err, "error")
                        return False
                except subprocess.TimeoutExpired:
                    self._wlog("❌ 更新超时（超过2分钟）", "error")
                    return False
                except Exception as e:
                    self._wlog("❌ 更新异常: " + str(e), "error")
                    return False

            try:
                subprocess.Popen([tortoise, "/command:commit", "/path:" + d])
                self._wlog("TortoiseSVN 提交对话框已打开: " + d, "ok")
            except Exception as e:
                self._wlog("TortoiseSVN 启动失败: " + str(e), "error")
                return False
        return True

    def _ulog(self, msg, level="info"):
        """上传SVN日志"""
        ts = datetime.now().strftime("%H:%M:%S")
        msg = f"[{ts}] {msg}"
        log_widget = self.upload_log_text
        def _update():
            try:
                if not log_widget.winfo_exists():
                    return
                log_widget.config(state="normal")
                log_widget.insert("end", msg + "\n", level)
                log_widget.see("end")
                log_widget.config(state="disabled")
            except Exception:
                pass
        self.root.after(0, _update)

    def _wf_execute_merge_table(self, step):
        import openpyxl

        input_paths = step.get("input_paths", [])
        target_dir = step.get("target_dir", "").strip()
        prefixes = step.get("merge_prefixes", [])

        if not input_paths:
            self._wlog("未配置输入文件/目录", "error")
            return False
        if not target_dir or not os.path.isdir(target_dir):
            self._wlog("目标文件夹无效: " + str(target_dir), "error")
            return False

        # 1. 收集所有输入 Excel/CSV 文件
        input_excel_files = []
        for p in input_paths:
            p = p.strip()
            if not os.path.exists(p):
                self._wlog("路径不存在: " + p, "warn")
                continue
            if os.path.isfile(p):
                ext = os.path.splitext(p)[1].lower()
                if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
                    input_excel_files.append(p)
            elif os.path.isdir(p):
                for root, dirs, files in os.walk(p):
                    for f in files:
                        ext = os.path.splitext(f)[1].lower()
                        if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
                            input_excel_files.append(os.path.join(root, f))

        if not input_excel_files:
            self._wlog("未找到任何 Excel 文件", "error")
            return False
        self._wlog("找到 " + str(len(input_excel_files)) + " 个输入 Excel 文件", "info")

        matched_pairs = []

        def _stem(name):
            return os.path.splitext(name)[0]

        EXCEL_EXTS = {".xlsx", ".xlsm", ".xls", ".csv"}

        def _find_excel_by_stem(stem_name, search_dir):
            for root, dirs, files in os.walk(search_dir):
                for f in files:
                    ext = os.path.splitext(f)[1].lower()
                    if ext in EXCEL_EXTS and _stem(f) == stem_name:
                        return os.path.join(root, f)
            return None

        def _find_excel_in_folder(stem_name, search_dir):
            for root, dirs, files in os.walk(search_dir):
                for d in dirs:
                    if d == stem_name:
                        folder_path = os.path.join(root, d)
                        for f in os.listdir(folder_path):
                            ext = os.path.splitext(f)[1].lower()
                            if ext in EXCEL_EXTS:
                                return os.path.join(folder_path, f)
            return None

        # 2. 为每个输入文件匹配目标文件
        for inp in input_excel_files:
            inp_name = os.path.basename(inp)
            inp_stem = _stem(inp_name)
            exact_match = None

            if prefixes:
                exact_match = _find_excel_by_stem(inp_stem, target_dir)
                if not exact_match:
                    exact_match = _find_excel_in_folder(inp_stem, target_dir)

                stripped = inp_name
                for pre in prefixes:
                    pre = pre.strip()
                    if pre and stripped.startswith(pre):
                        stripped = stripped[len(pre):]
                        break

                stripped_stem = _stem(stripped)

                if exact_match:
                    self._wlog("匹配: " + inp_name + "（原文件名直接匹配）", "info")
                else:
                    self._wlog("匹配: " + inp_name + "（去前缀: " + stripped + "）", "info")
                    exact_match = _find_excel_by_stem(stripped_stem, target_dir)
                    if not exact_match:
                        exact_match = _find_excel_in_folder(stripped_stem, target_dir)
            else:
                self._wlog("匹配: " + inp_name, "info")
                exact_match = _find_excel_by_stem(inp_stem, target_dir)
                if not exact_match:
                    exact_match = _find_excel_in_folder(inp_stem, target_dir)

            if exact_match:
                self._wlog("  -> 匹配到: " + exact_match, "ok")
                matched_pairs.append((inp, exact_match))
            else:
                display_name = inp_name
                if prefixes and _stem(stripped) != inp_stem:
                    display_name = display_name + "（去前缀后: " + stripped + "）"
                self._wlog("  -> 未匹配到目标文件: " + display_name, "warn")

        if not matched_pairs:
            self._wlog("没有可合并的匹配文件对", "error")
            return False

        # 3. 逐对处理
        total_updated = 0
        total_added = 0

        for inp_file, tgt_file in matched_pairs:
            self._wlog("处理: " + os.path.basename(inp_file) + " -> " + os.path.basename(tgt_file), "head")

            local_title_rows = int(step.get("title_rows", "1"))
            raw_id_col = step.get("id_col", "1").strip()
            try:
                local_id_col_num = int(raw_id_col)
            except ValueError:
                local_id_col_num = 1
            self._wlog("  标题行数: " + str(local_title_rows) + ", ID列号: " + str(local_id_col_num), "info")

            try:
                inp_wb = openpyxl.load_workbook(inp_file, data_only=True)
                tgt_wb = openpyxl.load_workbook(tgt_file)

                header_row = local_title_rows
                total_sheets = 0
                all_sheet_ops = []

                for inp_sn in inp_wb.sheetnames:
                    inp_ws = inp_wb[inp_sn]
                    if inp_ws.max_row < header_row + 1:
                        continue

                    sheet_col = None
                    for col in range(1, inp_ws.max_column + 1):
                        h = inp_ws.cell(row=header_row, column=col).value
                        if h and str(h).strip().lower() == "sheet":
                            sheet_col = col
                            break

                    if sheet_col is not None:
                        sheet_rows = {}
                        for row in range(header_row + 1, inp_ws.max_row + 1):
                            v = inp_ws.cell(row=row, column=sheet_col).value
                            if v is not None:
                                tgt_sn = str(v).strip()
                                if tgt_sn:
                                    sheet_rows.setdefault(tgt_sn, []).append(row)
                        for tgt_sn, rows in sheet_rows.items():
                            if tgt_sn not in tgt_wb.sheetnames:
                                self._wlog("  目标文件无 sheet: " + tgt_sn, "warn")
                                continue
                            tgt_ws = tgt_wb[tgt_sn]
                            for r in rows:
                                a, u, ops = self._merge_sheet_changes(
                                    inp_ws, tgt_ws, r, r, header_row, local_id_col_num)
                                total_added += a
                                total_updated += u
                                total_sheets += 1
                                if ops["updates"] or ops["inserts"]:
                                    ops["sheet"] = tgt_sn
                                    all_sheet_ops.append(ops)
                    else:
                        if inp_sn not in tgt_wb.sheetnames:
                            continue
                        tgt_ws = tgt_wb[inp_sn]
                        a, u, ops = self._merge_sheet_changes(
                            inp_ws, tgt_ws, header_row + 1, inp_ws.max_row,
                            header_row, local_id_col_num)
                        total_added += a
                        total_updated += u
                        total_sheets += 1
                        if ops["updates"] or ops["inserts"]:
                            ops["sheet"] = inp_sn
                            all_sheet_ops.append(ops)

                self._wlog("  处理了 " + str(total_sheets) + " 个 sheet", "info")
                inp_wb.close()
                tgt_wb.close()

                if all_sheet_ops:
                    if _check_office_lock(tgt_file):
                        self._wlog("  WPS/Excel 锁定文件存在，目标文件可能正在被编辑，跳过保存", "error")
                    else:
                        self._wlog("  Excel 后台写入中...", "info")
                        ok, err_msg = apply_via_excel(tgt_file, all_sheet_ops)
                        if ok:
                            self._wlog("  已保存: " + os.path.basename(tgt_file), "ok")
                        else:
                            self._wlog("  Excel 写入失败: " + err_msg, "error")
                else:
                    self._wlog("  无变更，跳过保存", "info")

            except Exception as e:
                self._wlog("  处理失败: " + str(e), "error")
                import traceback
                self._wlog("  " + traceback.format_exc(), "error")

        self._wlog("合并完成: 新增 " + str(total_added) + " 行, 更新 " + str(total_updated) + " 行", "ok")
        return True
